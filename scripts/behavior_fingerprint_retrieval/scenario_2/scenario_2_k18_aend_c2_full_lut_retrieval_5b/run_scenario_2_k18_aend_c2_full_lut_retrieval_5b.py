"""Run the fixed K18 Aend-LUT/C2-Query Full-425 retrieval task.

The task is deliberately a retrieval execution task, not a model-selection task.
The support, Ridge lambda, canonical preprocessing, Common-B probe, and
deterministic Top-1 tie rule are all validated before any state is fitted.
"""

# ruff: noqa: E402,E501

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from multiprocessing import get_context
from pathlib import Path
from typing import Any

for _name in (
    "OMP_NUM_THREADS",
    "OPENBLAS_NUM_THREADS",
    "MKL_NUM_THREADS",
    "NUMEXPR_NUM_THREADS",
    "VECLIB_MAXIMUM_THREADS",
):
    os.environ[_name] = "1"

import numpy as np
import pandas as pd

PROJECT_ROOT = next(
    parent
    for parent in Path(__file__).resolve().parents
    if (parent / "environment.yml").is_file() and (parent / "scripts").is_dir()
)
SCRIPTS_ROOT = PROJECT_ROOT / "scripts"
if str(SCRIPTS_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_ROOT))

from behavior_fingerprint_ranking_consistency.shared.distance_matrix import (  # noqa: E402
    compute_cnmse_distance_matrix,
)
from behavior_modeling.shared.basis_function_selection.envelope_dictionary import (  # noqa: E402
    DMAX,
    EnvelopeBasis,
    build_envelope_bank,
    build_envelope_dictionary,
)
from behavior_modeling.shared.basis_function_selection.model_solver import (  # noqa: E402
    fit_ridge,
)
from core.shared.metrics import nmse  # noqa: E402
from data_management.shared import build_state_table, load_by_id  # noqa: E402
from signal_segmentation.shared import (  # noqa: E402
    build_off_segments,
    build_partition_from_xin,
    get_common_probe,
    get_ilc_pair,
    preprocess_full_pair,
)

from behavior_fingerprint_retrieval.scenario_2.scenario_2_k18_aend_c2_full_lut_retrieval_5b.plot_scenario_2_k18_aend_c2_full_lut_retrieval_5b import (  # noqa: E402
    write_figures,
)

TASK_NAME = "scenario_2_k18_aend_c2_full_lut_retrieval_5b"
RESULT_ROOT = PROJECT_ROOT / "results" / "behavior_fingerprint_retrieval" / "scenario_2" /TASK_NAME
WORK_LOG = PROJECT_ROOT / "work_logs" / "behavior_fingerprint_retrieval" / "scenario_2" /TASK_NAME / "execution_log.txt"
HANDOFF_LOG = PROJECT_ROOT / "work_logs" / "core" / "codex_handoff" / "codex_handoff.txt"
FROZEN_ROOT = (
    PROJECT_ROOT
    / "results"
    / "retrieval_oriented_model_selection"
    / "scenario_2_self_first_joint_basis_ridge_lut_retrieval_search_5b"
)
FROZEN_SUMMARY_PATH = FROZEN_ROOT / "22_final_result_summary.json"
FROZEN_CHECKPOINT_PATH = FROZEN_ROOT / "23_checkpoint.json"
FROZEN_VALIDATION_PATH = FROZEN_ROOT / "24_validation_checks.json"
HISTORICAL_QUERY_PATH = FROZEN_ROOT / "final_single_add" / "query_metrics_long.csv"
HISTORICAL_DISTANCE_PATH = FROZEN_ROOT / "final_single_add" / "distance_matrices" / "candidate_0115.npy"

STATE_COUNT = 425
FULL_LENGTH = 24_576
ABC_LENGTHS = (12_288, 4_915, 7_373)
VALID_LENGTHS = (12_286, 4_913, 7_371)
COMMON_B_RAW_LENGTH = 4_915
COMMON_B_VALID_LENGTH = 4_913
COMMON_B_SOURCE_STATE = 0
COMMON_B_SHA256 = "b3a794e22b6beb7171f696d41685714da122f948ca8284553abaaa9a68c4abac"
RIDGE_LAMBDA = 3.16227766e-12
MODEL_CANDIDATE = "final_add_ENV_p04_m2_q1_lambda_3em12"
MODEL_K = 18
MODEL_P_MAX = 8
MODEL_SUPPORT_HASH = "263d4d30d09d651f9097fe86b50ad769f0b7e0cc7ac41c0d703f0a1dbf8b4f8c"
REAL_B_THRESHOLD_DB = -40.0
TIE_TOLERANCE_DB = 1e-12
WORKER_COUNT = 10
BLAS_THREADS_PER_WORKER = 1
EXPECTED_RAW_MANIFEST = "910b06e4be6b60b1f4dc63409a3083eb328765347fbec3b89a97a41d6633a1e0"
EXPECTED_RAW_FILE_COUNT = 429
EXPECTED_RAW_MAT_COUNT = 427
EXPECTED_RAW_BYTES = 2_258_448_137
EXPECTED_COUNTS = {"N_self": 270, "N_fallback": 149, "N_valid": 419, "N_fail": 6}

SUPPORT_IDS = (
    "LIN_d0",
    "LIN_d1",
    "LIN_d2",
    "ENV_p02_m0_q0",
    "ENV_p02_m1_q1",
    "ENV_p02_m2_q2",
    "ENV_p03_m0_q0",
    "ENV_p03_m0_q2",
    "ENV_p03_m1_q0",
    "ENV_p03_m1_q2",
    "ENV_p04_m0_q0",
    "ENV_p04_m0_q1",
    "ENV_p04_m2_q1",
    "ENV_p05_m0_q0",
    "ENV_p05_m2_q0",
    "ENV_p06_m0_q0",
    "ENV_p07_m0_q0",
    "ENV_p08_m0_q0",
)

_WORKER_TERMS: tuple[EnvelopeBasis, ...] | None = None
_WORKER_SUPPORT: tuple[int, ...] | None = None
_WORKER_COMMON_PHI: np.ndarray | None = None
_WORKER_STATE_TABLE: dict[int, dict[str, int | float]] | None = None


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%S%z")


def _json_default(value: Any) -> Any:
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer, int)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.ndarray):
        return value.tolist()
    raise TypeError(f"Cannot JSON serialize {type(value)!r}")


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=_json_default) + "\n",
        encoding="utf-8",
    )


def _append_log(message: str) -> None:
    WORK_LOG.parent.mkdir(parents=True, exist_ok=True)
    with WORK_LOG.open("a", encoding="utf-8") as handle:
        handle.write(f"{_now()} | {message.rstrip()}\n")


def _append_handoff(message: str) -> None:
    HANDOFF_LOG.parent.mkdir(parents=True, exist_ok=True)
    with HANDOFF_LOG.open("a", encoding="utf-8") as handle:
        handle.write(f"\n{_now()} | 新任务：{TASK_NAME}\n{message.rstrip()}\n")


def _vector(value: Any, name: str) -> np.ndarray:
    array = np.asarray(value)
    if array.ndim == 2 and 1 in array.shape:
        array = array.reshape(-1)
    if array.ndim != 1 or array.size != FULL_LENGTH or not np.iscomplexobj(array):
        raise ValueError(f"{name} must be a complex vector of length {FULL_LENGTH}, got {array.shape}")
    array = np.asarray(array, dtype=np.complex128)
    if not np.all(np.isfinite(array)):
        raise ValueError(f"{name} contains NaN or Inf")
    return array


def _scalar(value: Any, name: str) -> float:
    array = np.asarray(value).reshape(-1)
    if array.size != 1 or not np.isfinite(array[0]):
        raise ValueError(f"{name} must be one finite scalar")
    return float(array[0])


def _support_hash(basis_ids: tuple[str, ...] = SUPPORT_IDS) -> str:
    payload = json.dumps(list(basis_ids), ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def raw_manifest() -> dict[str, object]:
    """Hash raw relative names and sizes without reading or changing raw data."""

    raw_root = PROJECT_ROOT / "data" / "raw"
    digest = hashlib.sha256()
    files = sorted(
        (file for file in raw_root.rglob("*") if file.is_file()),
        key=lambda item: str(item).lower(),
    )
    total_bytes = 0
    mat_count = 0
    for file in files:
        digest.update(file.relative_to(raw_root).as_posix().encode("utf-8") + b"\0")
        size = int(file.stat().st_size)
        digest.update(size.to_bytes(8, "little"))
        total_bytes += size
        mat_count += int(file.suffix.lower() == ".mat")
    return {
        "sha256": digest.hexdigest(),
        "file_count": len(files),
        "mat_count": mat_count,
        "bytes": total_bytes,
    }


def _raw_manifest_gate() -> dict[str, object]:
    value = raw_manifest()
    expected = {
        "sha256": EXPECTED_RAW_MANIFEST,
        "file_count": EXPECTED_RAW_FILE_COUNT,
        "mat_count": EXPECTED_RAW_MAT_COUNT,
        "bytes": EXPECTED_RAW_BYTES,
    }
    if value != expected:
        raise RuntimeError(f"data/raw manifest differs from frozen baseline: {value}")
    return value


def _parse_json_field(value: Any, name: str) -> list[Any]:
    if isinstance(value, str):
        parsed = json.loads(value)
    else:
        parsed = value
    if not isinstance(parsed, list):
        raise RuntimeError(f"frozen {name} must be a list")
    return parsed


def load_frozen_model_contract() -> dict[str, Any]:
    """Validate the Self-First final champion without importing its task code."""

    for path in (FROZEN_SUMMARY_PATH, FROZEN_CHECKPOINT_PATH, FROZEN_VALIDATION_PATH):
        if not path.is_file():
            raise FileNotFoundError(f"required frozen model file is missing: {path}")
    summary = json.loads(FROZEN_SUMMARY_PATH.read_text(encoding="utf-8"))
    checkpoint = json.loads(FROZEN_CHECKPOINT_PATH.read_text(encoding="utf-8"))
    validation = json.loads(FROZEN_VALIDATION_PATH.read_text(encoding="utf-8"))
    candidate = summary.get("best_candidate")
    checkpoint_candidate = checkpoint.get("best_candidate")
    if not isinstance(candidate, dict) or not isinstance(checkpoint_candidate, dict):
        raise RuntimeError("frozen result lacks best_candidate contract")
    if summary.get("status") != "SUCCESS" or checkpoint.get("status") != "SUCCESS":
        raise RuntimeError("frozen Self-First result is not SUCCESS")
    if checkpoint.get("phase") != "completed":
        raise RuntimeError("frozen Self-First checkpoint is not completed")
    if summary.get("objective_version") != "SELF_FIRST_V1":
        raise RuntimeError("unexpected frozen objective version")
    basis_ids = tuple(str(item) for item in _parse_json_field(candidate.get("basis_ids"), "basis_ids"))
    support_indices = tuple(int(item) for item in _parse_json_field(candidate.get("support_indices"), "support_indices"))
    expected_terms = {term.basis_id: term for term in build_envelope_dictionary()}
    expected_indices = tuple(expected_terms[item].index for item in SUPPORT_IDS)
    checks = {
        "candidate": candidate.get("candidate_name") == MODEL_CANDIDATE,
        "K": int(candidate.get("K", -1)) == MODEL_K,
        "lambda": math.isclose(float(candidate.get("lambda")), RIDGE_LAMBDA, rel_tol=0.0, abs_tol=1e-24),
        "support_hash": candidate.get("support_hash") == MODEL_SUPPORT_HASH,
        "basis_ids": basis_ids == SUPPORT_IDS,
        "support_indices": support_indices == expected_indices,
        "checkpoint_candidate": checkpoint_candidate.get("candidate_name") == MODEL_CANDIDATE,
        "checkpoint_support_hash": checkpoint_candidate.get("support_hash") == MODEL_SUPPORT_HASH,
        "checkpoint_lambda": math.isclose(float(checkpoint_candidate.get("lambda")), RIDGE_LAMBDA, rel_tol=0.0, abs_tol=1e-24),
        "checkpoint_state_count": int(checkpoint.get("state_count", -1)) == STATE_COUNT,
        "historical_counts": all(
            int(summary.get(f"best_{key}", -1)) == value for key, value in EXPECTED_COUNTS.items()
        ),
    }
    if not all(checks.values()):
        raise RuntimeError(f"frozen K18 contract failed: {checks}")
    return {
        "candidate": MODEL_CANDIDATE,
        "K": MODEL_K,
        "lambda": RIDGE_LAMBDA,
        "dmax": DMAX,
        "p_max": MODEL_P_MAX,
        "support_hash": MODEL_SUPPORT_HASH,
        "support_ids": list(SUPPORT_IDS),
        "support_indices": list(expected_indices),
        "objective_version": summary["objective_version"],
        "historical_counts": {
            key: int(summary[f"best_{key}"]) for key in EXPECTED_COUNTS
        },
        "nested_cv_status": summary.get("nested_cv_status"),
        "source_summary": str(FROZEN_SUMMARY_PATH),
        "source_checkpoint": str(FROZEN_CHECKPOINT_PATH),
        "source_validation": str(FROZEN_VALIDATION_PATH),
        "validation_class_identity": validation.get("class_identity", []),
        "contract_checks": checks,
    }


def _build_support_contract() -> tuple[tuple[EnvelopeBasis, ...], tuple[int, ...]]:
    terms = tuple(build_envelope_dictionary())
    by_id = {term.basis_id: term for term in terms}
    support_indices = tuple(int(by_id[item].index) for item in SUPPORT_IDS)
    if len(terms) != 75 or len(support_indices) != MODEL_K:
        raise RuntimeError("Envelope75/K18 dictionary contract failed")
    if _support_hash() != MODEL_SUPPORT_HASH:
        raise RuntimeError("K18 support hash calculation changed")
    if any(terms[index].basis_id != basis_id for index, basis_id in zip(support_indices, SUPPORT_IDS, strict=True)):
        raise RuntimeError("K18 support ordering changed")
    return terms, support_indices


def build_common_b_probe(
    terms: tuple[EnvelopeBasis, ...],
    support_indices: tuple[int, ...],
) -> tuple[np.ndarray, np.ndarray, dict[str, object]]:
    data = load_by_id(COMMON_B_SOURCE_STATE)
    xin = _vector(data["xin"], "xin")
    partition = build_partition_from_xin(xin)
    if (partition.n_a, partition.n_b, partition.n_c) != ABC_LENGTHS:
        raise RuntimeError("canonical ABC ownership lengths changed")
    common_b = np.asarray(get_common_probe(data, partition), dtype=np.complex128)
    if common_b.shape != (COMMON_B_RAW_LENGTH,) or not np.all(np.isfinite(common_b)):
        raise RuntimeError(f"Common-B shape/finite gate failed: {common_b.shape}")
    phi_b = np.asarray(build_envelope_bank(common_b, terms)[:, support_indices], dtype=np.complex128)
    if phi_b.shape != (COMMON_B_VALID_LENGTH, MODEL_K) or not np.all(np.isfinite(phi_b)):
        raise RuntimeError(f"Common-B design shape/finite gate failed: {phi_b.shape}")
    digest = hashlib.sha256(np.ascontiguousarray(common_b).tobytes()).hexdigest()
    if digest != COMMON_B_SHA256:
        raise RuntimeError(f"Common-B SHA changed: {digest}")
    metadata = {
        "source_state": COMMON_B_SOURCE_STATE,
        "source": "State 0 canonical xin B via signal_segmentation.shared.get_common_probe",
        "raw_length": COMMON_B_RAW_LENGTH,
        "valid_length": COMMON_B_VALID_LENGTH,
        "dtype": str(common_b.dtype),
        "sha256": digest,
        "finite": bool(np.all(np.isfinite(common_b))),
        "peak": float(np.max(np.abs(common_b))),
        "rms": float(np.sqrt(np.mean(np.abs(common_b) ** 2))),
        "ABC_lengths": list(ABC_LENGTHS),
        "dmax": DMAX,
    }
    return common_b, phi_b, metadata


def _fit_side(
    x_segment: np.ndarray,
    y_segment: np.ndarray,
    terms: tuple[EnvelopeBasis, ...],
    support_indices: tuple[int, ...],
) -> tuple[dict[str, object], np.ndarray]:
    full_bank = build_envelope_bank(x_segment, terms)
    phi = np.asarray(full_bank[:, support_indices], dtype=np.complex128)
    target = np.asarray(y_segment[DMAX:], dtype=np.complex128)
    if phi.shape != (x_segment.size - DMAX, MODEL_K) or target.shape != (x_segment.size - DMAX,):
        raise RuntimeError(f"fit shape mismatch: phi={phi.shape}, target={target.shape}")
    fitted = fit_ridge(phi, target, RIDGE_LAMBDA)
    finite = bool(
        np.all(np.isfinite(phi))
        and np.all(np.isfinite(target))
        and np.all(np.isfinite(fitted.theta))
        and np.all(np.isfinite(fitted.prediction))
        and np.isfinite(fitted.condition_number)
    )
    return {
        "train_NMSE_dB": float(fitted.nmse_db),
        "rank": int(fitted.rank),
        "condition_number": float(fitted.condition_number),
        "coefficient_norm": float(np.linalg.norm(fitted.theta)),
        "finite": finite,
    }, np.asarray(fitted.theta, dtype=np.complex128)


def _worker_init(support_indices: tuple[int, ...], common_phi: np.ndarray) -> None:
    global _WORKER_TERMS, _WORKER_SUPPORT, _WORKER_COMMON_PHI, _WORKER_STATE_TABLE
    _WORKER_TERMS = tuple(build_envelope_dictionary())
    _WORKER_SUPPORT = tuple(int(item) for item in support_indices)
    _WORKER_COMMON_PHI = np.asarray(common_phi, dtype=np.complex128)
    _WORKER_STATE_TABLE = {int(row["state_id"]): row for row in build_state_table()}
    if _WORKER_COMMON_PHI.shape != (COMMON_B_VALID_LENGTH, MODEL_K):
        raise RuntimeError("worker Common-B design shape changed")


def state_worker(state_id: int) -> dict[str, object]:
    if (
        _WORKER_TERMS is None
        or _WORKER_SUPPORT is None
        or _WORKER_COMMON_PHI is None
        or _WORKER_STATE_TABLE is None
    ):
        raise RuntimeError("K18 worker is not initialized")
    normalized_id = int(state_id)
    data = load_by_id(normalized_id)
    xin = _vector(data["xin"], "xin")
    if _scalar(data["freqSample_Hz"], "freqSample_Hz") != 100_000_000.0:
        raise RuntimeError(f"state {normalized_id} is not Fs=100 MHz")
    if _scalar(data["bandWidth_MHz"], "bandWidth_MHz") != 20.0:
        raise RuntimeError(f"state {normalized_id} is not B=20 MHz")
    partition = build_partition_from_xin(xin)
    if (partition.n_a, partition.n_b, partition.n_c) != ABC_LENGTHS:
        raise RuntimeError(f"state {normalized_id} ABC lengths changed")
    histories = np.asarray(data["xin_pd_ori_ilc"])
    outputs = np.asarray(data["yout_withdpd_ori_ilc"])
    if histories.ndim != 2 or outputs.shape != histories.shape or histories.shape[1] < 2:
        raise RuntimeError(f"state {normalized_id} lacks a valid C2 column")
    aend_index = int(histories.shape[1] - 1)
    c2_index = 1
    aend_pair = get_ilc_pair(data, aend_index)
    c2_pair = get_ilc_pair(data, c2_index)
    aend = preprocess_full_pair(
        aend_pair.input_full,
        aend_pair.output_raw_full,
        partition,
        pair_type=aend_pair.pair_type,
        iteration_index=aend_pair.iteration_index,
        input_peak_normalization_factor=aend_pair.input_peak_normalization_factor,
    )
    c2 = preprocess_full_pair(
        c2_pair.input_full,
        c2_pair.output_raw_full,
        partition,
        pair_type=c2_pair.pair_type,
        iteration_index=c2_pair.iteration_index,
        input_peak_normalization_factor=c2_pair.input_peak_normalization_factor,
    )
    off = build_off_segments(data, partition)
    a_metrics, theta_a = _fit_side(aend["A"].input, aend["A"].output, _WORKER_TERMS, _WORKER_SUPPORT)
    c_metrics, theta_c = _fit_side(c2["C"].input, c2["C"].output, _WORKER_TERMS, _WORKER_SUPPORT)
    a_b_phi = np.asarray(build_envelope_bank(aend["B"].input, _WORKER_TERMS)[:, _WORKER_SUPPORT], dtype=np.complex128)
    c_b_phi = np.asarray(build_envelope_bank(c2["B"].input, _WORKER_TERMS)[:, _WORKER_SUPPORT], dtype=np.complex128)
    a_b_prediction = a_b_phi @ theta_a
    c_b_prediction = c_b_phi @ theta_c
    y_a_b = np.asarray(aend["B"].output[DMAX:], dtype=np.complex128)
    y_c_b = np.asarray(c2["B"].output[DMAX:], dtype=np.complex128)
    real_b = np.asarray(off["B"].output[DMAX:], dtype=np.complex128)
    if real_b.shape != (COMMON_B_VALID_LENGTH,) or not np.all(np.isfinite(real_b)):
        raise RuntimeError(f"state {normalized_id} Real-B shape/finite gate failed")
    state = _WORKER_STATE_TABLE[normalized_id]
    low_acpr = _scalar(data["acpr_low_withoutdpd"], "acpr_low_withoutdpd")
    high_acpr = _scalar(data["acpr_upper_withoutdpd"], "acpr_upper_withoutdpd")
    return {
        "state_R": normalized_id,
        "funMng": int(state["funMng"]),
        "funAng": int(state["funAng"]),
        "secMng": int(state["secMng"]),
        "secAng": int(state["secAng"]),
        "Vm": float(state["Vm"]),
        "Pin": float(state["Pin"]),
        "nmse_withoutdpd_dB": _scalar(data["nmse_withoutdpd"], "nmse_withoutdpd"),
        "acpr_low_withoutdpd_dBc": low_acpr,
        "acpr_high_withoutdpd_dBc": high_acpr,
        "acpr_withoutdpd_mean_dBc": (low_acpr + high_acpr) / 2.0,
        "ilc_A_end": aend_index + 1,
        "ilc_C_2": c2_index + 1,
        "Y_Aend_train_NMSE_dB": float(a_metrics["train_NMSE_dB"]),
        "Y_Aend_B_NMSE_dB": float(nmse(y_a_b, a_b_prediction)),
        "Y_Aend_rank": int(a_metrics["rank"]),
        "Y_Aend_condition_number": float(a_metrics["condition_number"]),
        "Y_Aend_coefficient_norm": float(a_metrics["coefficient_norm"]),
        "Y_Aend_finite": bool(a_metrics["finite"]),
        "Y_C2_train_NMSE_dB": float(c_metrics["train_NMSE_dB"]),
        "Y_C2_B_NMSE_dB": float(nmse(y_c_b, c_b_prediction)),
        "Y_C2_rank": int(c_metrics["rank"]),
        "Y_C2_condition_number": float(c_metrics["condition_number"]),
        "Y_C2_coefficient_norm": float(c_metrics["coefficient_norm"]),
        "Y_C2_finite": bool(c_metrics["finite"]),
        "theta_Aend": theta_a,
        "theta_C2": theta_c,
        "LUT_fingerprint": np.asarray(_WORKER_COMMON_PHI @ theta_a, dtype=np.complex128),
        "Query_fingerprint": np.asarray(_WORKER_COMMON_PHI @ theta_c, dtype=np.complex128),
        "real_B_valid": real_b,
    }


def _top1(distance: np.ndarray) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    if distance.shape != (STATE_COUNT, STATE_COUNT):
        raise RuntimeError(f"distance matrix shape changed: {distance.shape}")
    if np.isnan(distance).any() or np.isposinf(distance).any():
        raise RuntimeError("fingerprint distance contains NaN or +Inf")
    state_ids = np.arange(STATE_COUNT, dtype=np.int64)
    selected = np.empty(STATE_COUNT, dtype=np.int64)
    selected_distance = np.empty(STATE_COUNT, dtype=np.float64)
    true_rank = np.empty(STATE_COUNT, dtype=np.int64)
    tie_count = np.empty(STATE_COUNT, dtype=np.int64)
    for state_id in range(STATE_COUNT):
        row = np.asarray(distance[state_id], dtype=np.float64)
        minimum = float(np.min(row))
        tied = np.flatnonzero(row <= minimum + TIE_TOLERANCE_DB)
        selected[state_id] = int(np.min(tied))
        selected_distance[state_id] = float(row[selected[state_id]])
        tie_count[state_id] = int(tied.size)
        order = np.lexsort((state_ids, row))
        true_rank[state_id] = int(np.flatnonzero(order == state_id)[0] + 1)
    return selected, selected_distance, true_rank, tie_count


def _finite_stats(values: np.ndarray) -> dict[str, object]:
    finite = np.asarray(values, dtype=float)
    finite = finite[np.isfinite(finite)]
    if finite.size == 0:
        return {"count": 0, "median_dB": None, "mean_dB": None, "Q05_dB": None, "Q95_dB": None}
    return {
        "count": int(finite.size),
        "median_dB": float(np.median(finite)),
        "mean_dB": float(np.mean(finite)),
        "Q05_dB": float(np.quantile(finite, 0.05)),
        "Q95_dB": float(np.quantile(finite, 0.95)),
        "min_dB": float(np.min(finite)),
        "max_dB": float(np.max(finite)),
    }


def _load_historical_regression(
    distance: np.ndarray,
    selected: np.ndarray,
    real_b_values: np.ndarray,
    classes: np.ndarray,
) -> dict[str, object]:
    result: dict[str, object] = {
        "reference_query_metrics": str(HISTORICAL_QUERY_PATH),
        "reference_distance_matrix": str(HISTORICAL_DISTANCE_PATH),
        "query_state_match": False,
        "class_match": False,
        "real_B_metric_max_abs_difference_dB": None,
        "distance_matrix_shape_match": False,
        "distance_matrix_max_abs_difference": None,
        "distance_matrix_match_at_1e-9": False,
    }
    if not HISTORICAL_QUERY_PATH.is_file() or not HISTORICAL_DISTANCE_PATH.is_file():
        result["note"] = "historical exact candidate artifacts are unavailable"
        return result
    history = pd.read_csv(HISTORICAL_QUERY_PATH)
    history = history.loc[history["candidate_id"].astype(int) == 115].sort_values("State_R")
    history_distance = np.asarray(np.load(HISTORICAL_DISTANCE_PATH, allow_pickle=False), dtype=np.float64)
    if len(history) != STATE_COUNT:
        raise RuntimeError(f"historical candidate 115 row count changed: {len(history)}")
    if history_distance.shape != distance.shape:
        raise RuntimeError(f"historical candidate distance shape changed: {history_distance.shape}")
    historical_q = history["State_Q"].to_numpy(dtype=np.int64)
    historical_class = history["retrieval_class"].astype(str).replace({"SHAREABLE_FALLBACK": "FALLBACK"}).to_numpy(dtype=object)
    current_class = np.asarray(classes, dtype=object)
    history_real = history["retrieved_real_B_CNMSE_dB"].to_numpy(dtype=float)
    finite_mask = np.isfinite(history_real) & np.isfinite(real_b_values)
    real_error = float(np.max(np.abs(history_real[finite_mask] - real_b_values[finite_mask]))) if np.any(finite_mask) else 0.0
    distance_error = float(np.max(np.abs(distance - history_distance)))
    result.update(
        {
            "query_state_match": bool(np.array_equal(selected, historical_q)),
            "class_match": bool(np.array_equal(current_class, historical_class)),
            "real_B_metric_max_abs_difference_dB": real_error,
            "historical_real_B_csv_precision_dB": 6,
            "distance_matrix_shape_match": True,
            "distance_matrix_max_abs_difference": distance_error,
            "distance_matrix_match_at_1e-9": bool(distance_error <= 1e-9),
            "distance_matrix_match_at_1e-8": bool(distance_error <= 1e-8),
        }
    )
    return result


def _basis_definition(terms: tuple[EnvelopeBasis, ...], support_indices: tuple[int, ...]) -> list[dict[str, object]]:
    output: list[dict[str, object]] = []
    for support_index, global_index in enumerate(support_indices):
        term = terms[global_index]
        output.append(
            {
                "basis_index": support_index,
                "global_dictionary_index": int(global_index),
                "basis_name": term.basis_id,
                "p": int(term.order),
                "m": int(term.signal_delay),
                "q": None if term.envelope_delay is None else int(term.envelope_delay),
                "type": "LINEAR" if term.order == 1 else "ENVELOPE",
                "formula": term.formula,
            }
        )
    return output


def _load_config(row: dict[str, object]) -> str:
    return (
        f"funMng={int(row['funMng'])}, funAng={int(row['funAng'])}, "
        f"secMng={int(row['secMng'])}, secAng={int(row['secAng'])}"
    )


def _write_excel(
    frame: pd.DataFrame,
    summary: dict[str, object],
    model_definition: dict[str, object],
    path: Path,
) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "summary"
    summary_items = [(str(key), value) for key, value in summary.items()]
    summary_sheet.append(["field", "value"])
    for key, value in summary_items:
        encoded = json.dumps(value, ensure_ascii=False, default=_json_default) if isinstance(value, (dict, list, tuple)) else value
        summary_sheet.append([key, encoded])

    state_sheet = workbook.create_sheet("state_results")
    excel_frame = frame.copy()
    if "retrieved_real_B_CNMSE_dB" in excel_frame:
        excel_frame["retrieved_real_B_CNMSE_dB"] = [
            "-Inf" if np.isneginf(value) else float(value)
            for value in excel_frame["retrieved_real_B_CNMSE_dB"].to_numpy(dtype=float)
        ]
    state_sheet.append(list(excel_frame.columns))
    for row in excel_frame.itertuples(index=False, name=None):
        state_sheet.append(list(row))

    model_sheet = workbook.create_sheet("model_definition")
    model_columns = [
        "basis_index",
        "basis_name",
        "p",
        "m",
        "q",
        "type",
        "formula",
        "global_dictionary_index",
        "candidate",
        "K",
        "lambda",
        "dmax",
        "support_hash",
    ]
    model_sheet.append(model_columns)
    for basis in model_definition["support"]:
        model_sheet.append(
            [
                basis["basis_index"],
                basis["basis_name"],
                basis["p"],
                basis["m"],
                basis["q"],
                basis["type"],
                basis["formula"],
                basis["global_dictionary_index"],
                model_definition["candidate"],
                model_definition["K"],
                model_definition["lambda"],
                model_definition["dmax"],
                model_definition["support_hash"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for column_index in range(1, sheet.max_column + 1):
            values = [sheet.cell(row=row_index, column=column_index).value for row_index in range(1, min(sheet.max_row, 51) + 1)]
            width = min(max(12, max(len(str(value)) for value in values) + 2), 42)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (float, np.floating)):
                    cell.number_format = "0.000000"
    workbook.save(path)
    loaded = load_workbook(path, read_only=True, data_only=False)
    try:
        if loaded["state_results"].max_row != STATE_COUNT + 1:
            raise RuntimeError("Excel state_results row count is not 425")
        if loaded["state_results"].max_column != len(frame.columns):
            raise RuntimeError("Excel state_results column count differs from CSV")
    finally:
        loaded.close()


def _write_task_definition(contract: dict[str, Any], common_meta: dict[str, object], raw_before: dict[str, object]) -> None:
    lines = [
        f"Task: {TASK_NAME}",
        "Mode: new task; fixed-model Full-425 LUT retrieval.",
        "Scientific definition: fixed K18 model; Aend DPD-on model outputs form the LUT; C2 DPD-on model outputs form online queries; both use the same Common-B probe.",
        f"Candidate: {MODEL_CANDIDATE}",
        f"K: {MODEL_K}",
        f"Ridge lambda: {RIDGE_LAMBDA:.12g}",
        f"dmax: {DMAX}; p_max: {MODEL_P_MAX}",
        f"Support hash: {MODEL_SUPPORT_HASH}",
        f"Support IDs: {list(SUPPORT_IDS)}",
        "Aend: last valid ILC input/output column, represented by aend_index=history_columns-1.",
        "C2: literal second ILC column, zero-based index 1.",
        f"Scenario 2 states: 0...{STATE_COUNT - 1}; allow_self=true.",
        f"ABC lengths: {ABC_LENGTHS}; valid lengths after dmax: {VALID_LENGTHS}.",
        f"Common-B metadata: {json.dumps(common_meta, ensure_ascii=False, sort_keys=True)}",
        f"Raw manifest before: {json.dumps(raw_before, ensure_ascii=False, sort_keys=True)}",
        "Solver: behavior_modeling.shared.basis_function_selection.model_solver.fit_ridge; augmented complex least-squares convention.",
        "Excluded: basis search, Forward, Backward, Beam, Floating, Swap, Ridge scan, Dense Ridge, LOO, Nested CV, Self-First rerun, lambda optimization.",
    ]
    (RESULT_ROOT / "00_task_definition.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_checkpoint(phase: str, **payload: object) -> None:
    _write_json(
        RESULT_ROOT / "09_checkpoint.json",
        {
            "task_name": TASK_NAME,
            "phase": phase,
            "state_count": STATE_COUNT,
            "worker_count": WORKER_COUNT,
            "blas_threads_per_worker": BLAS_THREADS_PER_WORKER,
            "allow_self": True,
            "Aend_definition": "last_valid_ilc_column",
            "C2_definition": "literal_second_ilc_column",
            "dmax": DMAX,
            "K": MODEL_K,
            "lambda": RIDGE_LAMBDA,
            "candidate": MODEL_CANDIDATE,
            "support_hash": MODEL_SUPPORT_HASH,
            "common_B_sha256": COMMON_B_SHA256,
            **payload,
        },
    )


def run(*, allow_existing: bool = False) -> dict[str, object]:
    if RESULT_ROOT.exists() and any(RESULT_ROOT.iterdir()) and not allow_existing:
        raise RuntimeError(
            f"result directory is not empty; refusing to overwrite without --allow-existing: {RESULT_ROOT}"
        )
    RESULT_ROOT.mkdir(parents=True, exist_ok=True)
    WORK_LOG.parent.mkdir(parents=True, exist_ok=True)
    _append_log("START fixed K18 Aend/C2 Full-425 LUT retrieval")
    raw_before = _raw_manifest_gate()
    frozen_contract = load_frozen_model_contract()
    terms, support_indices = _build_support_contract()
    common_b, common_phi, common_meta = build_common_b_probe(terms, support_indices)
    _write_task_definition(frozen_contract, common_meta, raw_before)
    model_definition = {
        "candidate": MODEL_CANDIDATE,
        "K": MODEL_K,
        "lambda": RIDGE_LAMBDA,
        "dmax": DMAX,
        "p_max": MODEL_P_MAX,
        "support_hash": MODEL_SUPPORT_HASH,
        "support": _basis_definition(terms, support_indices),
        "basis_definition": "phi[p,m,q](n)=x[n-m] * abs(x[n-q])**(p-1)",
        "source_frozen_summary": str(FROZEN_SUMMARY_PATH),
    }
    _write_json(RESULT_ROOT / "00_preflight_contract.json", {"frozen_model": frozen_contract, "common_B": common_meta, "raw_manifest_before": raw_before})
    _write_json(RESULT_ROOT / "03_model_definition.json", model_definition)
    _write_json(RESULT_ROOT / "00_common_B_metadata.json", common_meta)
    _write_checkpoint("preflight_complete", completed_model_state_ids=[])
    _append_log(f"PRECHECK PASS raw={raw_before}; candidate={MODEL_CANDIDATE}; K={MODEL_K}; lambda={RIDGE_LAMBDA:.12g}; Common-B={COMMON_B_SHA256}")

    start = time.time()
    model_results: list[dict[str, object]] = []
    context = get_context("spawn")
    with ProcessPoolExecutor(
        max_workers=WORKER_COUNT,
        mp_context=context,
        initializer=_worker_init,
        initargs=(support_indices, common_phi),
    ) as executor:
        futures = {executor.submit(state_worker, state_id): state_id for state_id in range(STATE_COUNT)}
        for completed, future in enumerate(as_completed(futures), start=1):
            state_id = futures[future]
            try:
                model_results.append(future.result())
            except Exception as exc:
                for pending in futures:
                    pending.cancel()
                raise RuntimeError(f"state {state_id} fixed K18 evaluation failed") from exc
            if completed == 1 or completed % 25 == 0 or completed == STATE_COUNT:
                print(f"fixed K18 model states {completed}/{STATE_COUNT}", flush=True)
    model_results.sort(key=lambda row: int(row["state_R"]))
    if [int(row["state_R"]) for row in model_results] != list(range(STATE_COUNT)):
        raise RuntimeError("model result state order/count is not 0...424")
    lut_fingerprints = np.stack([row["LUT_fingerprint"] for row in model_results]).astype(np.complex128)
    query_fingerprints = np.stack([row["Query_fingerprint"] for row in model_results]).astype(np.complex128)
    real_b = np.stack([row["real_B_valid"] for row in model_results]).astype(np.complex128)
    theta_a = np.stack([row["theta_Aend"] for row in model_results]).astype(np.complex128)
    theta_c = np.stack([row["theta_C2"] for row in model_results]).astype(np.complex128)
    for matrix, name, shape in (
        (lut_fingerprints, "LUT fingerprints", (STATE_COUNT, COMMON_B_VALID_LENGTH)),
        (query_fingerprints, "Query fingerprints", (STATE_COUNT, COMMON_B_VALID_LENGTH)),
        (real_b, "Real-B", (STATE_COUNT, COMMON_B_VALID_LENGTH)),
        (theta_a, "Aend coefficients", (STATE_COUNT, MODEL_K)),
        (theta_c, "C2 coefficients", (STATE_COUNT, MODEL_K)),
    ):
        if matrix.shape != shape or not np.all(np.isfinite(matrix)):
            raise RuntimeError(f"{name} shape/finite gate failed: {matrix.shape}")
    _write_checkpoint("model_fitting_complete", completed_model_state_ids=list(range(STATE_COUNT)))
    _append_log(f"MODEL FITTING COMPLETE states={STATE_COUNT}; elapsed_s={time.time() - start:.1f}")

    distance = compute_cnmse_distance_matrix(query_fingerprints, lut_fingerprints)
    real_b_distance = compute_cnmse_distance_matrix(real_b, real_b)
    selected, selected_distance, true_rank, tie_count = _top1(distance)
    state_ids = np.arange(STATE_COUNT, dtype=np.int64)
    retrieved_real_b = real_b_distance[state_ids, selected]
    self_hit = selected == state_ids
    real_b_shareable = retrieved_real_b < REAL_B_THRESHOLD_DB
    fallback = (~self_hit) & real_b_shareable
    fail = (~self_hit) & (~real_b_shareable)
    retrieval_class = np.full(STATE_COUNT, "FAIL", dtype=object)
    retrieval_class[self_hit] = "SELF"
    retrieval_class[fallback] = "FALLBACK"
    counts = {
        "N_self": int(np.count_nonzero(self_hit)),
        "N_fallback": int(np.count_nonzero(fallback)),
        "N_valid": int(np.count_nonzero(self_hit | fallback)),
        "N_fail": int(np.count_nonzero(fail)),
    }
    abs_delta = np.abs(selected - state_ids).astype(np.int64)
    rows: list[dict[str, object]] = []
    state_table = build_state_table()
    for index, row in enumerate(model_results):
        state = state_table[index]
        retrieved = state_table[int(selected[index])]
        row_copy = {
            key: value
            for key, value in row.items()
            if key not in {"theta_Aend", "theta_C2", "LUT_fingerprint", "Query_fingerprint", "real_B_valid"}
        }
        row_copy.update(
            {
                "state_R": index,
                "load_config": _load_config(row_copy),
                "funAng_deg": int(state["funAng"]),
                "secAng_deg": int(state["secAng"]),
                "retrieved_state_Q": int(selected[index]),
                "retrieved_load_config": _load_config(retrieved),
                "retrieved_funMng": int(retrieved["funMng"]),
                "retrieved_funAng_deg": int(retrieved["funAng"]),
                "retrieved_secMng": int(retrieved["secMng"]),
                "retrieved_secAng_deg": int(retrieved["secAng"]),
                "retrieved_state_delta": int(selected[index] - index),
                "retrieved_state_abs_delta": int(abs_delta[index]),
                "retrieval_fingerprint_CNMSE_dB": float(selected_distance[index]),
                "retrieved_real_B_CNMSE_dB": float(retrieved_real_b[index]),
                "self_hit": bool(self_hit[index]),
                "real_B_shareable_lt_minus40": bool(real_b_shareable[index]),
                "retrieval_class": str(retrieval_class[index]),
                "true_state_rank": int(true_rank[index]),
                "minimum_tie_count": int(tie_count[index]),
            }
        )
        rows.append(row_copy)
    state_frame = pd.DataFrame(rows).sort_values("state_R").reset_index(drop=True)
    state_frame["retrieved_real_B_CNMSE_plot_dB"] = np.where(
        np.isfinite(state_frame["retrieved_real_B_CNMSE_dB"].to_numpy(dtype=float)),
        state_frame["retrieved_real_B_CNMSE_dB"].to_numpy(dtype=float),
        float(np.min(state_frame.loc[np.isfinite(state_frame["retrieved_real_B_CNMSE_dB"]), "retrieved_real_B_CNMSE_dB"]) - 5.0),
    )
    state_frame.to_csv(RESULT_ROOT / "01_state_results.csv", index=False)
    np.save(RESULT_ROOT / "04_fingerprint_distance_matrix.npy", distance)
    np.savez_compressed(
        RESULT_ROOT / "05_lut_fingerprints.npz",
        state_ids=state_ids,
        fingerprints=lut_fingerprints,
        common_B=common_b,
    )
    np.savez_compressed(
        RESULT_ROOT / "06_query_fingerprints.npz",
        state_ids=state_ids,
        fingerprints=query_fingerprints,
        common_B=common_b,
    )
    np.savez_compressed(
        RESULT_ROOT / "10_model_coefficients.npz",
        state_ids=state_ids,
        theta_Aend=theta_a,
        theta_C2=theta_c,
    )
    _write_checkpoint("retrieval_complete", completed_model_state_ids=list(range(STATE_COUNT)), retrieval_rows=STATE_COUNT, **counts)
    _append_log(f"RETRIEVAL COMPLETE counts={counts}; matrix={distance.shape}; Real-B strict threshold={REAL_B_THRESHOLD_DB} dB")

    historical = _load_historical_regression(distance, selected, retrieved_real_b, retrieval_class)
    finite_real_stats = _finite_stats(retrieved_real_b)
    delta_stats = {
        "mean_abs_state_delta": float(np.mean(abs_delta)),
        "median_abs_state_delta": float(np.median(abs_delta)),
        "max_abs_state_delta": int(np.max(abs_delta)),
    }
    figure_metadata = write_figures(state_frame)
    raw_after = _raw_manifest_gate()
    summary: dict[str, object] = {
        "status": "SUCCESS"
        if counts == EXPECTED_COUNTS
        and bool(historical.get("query_state_match", False))
        and bool(historical.get("class_match", False))
        and bool(historical.get("distance_matrix_match_at_1e-8", False))
        else "REGRESSION_MISMATCH",
        "task_name": TASK_NAME,
        "module_name": "behavior_fingerprint_retrieval",
        "candidate": MODEL_CANDIDATE,
        "K": MODEL_K,
        "lambda": RIDGE_LAMBDA,
        "dmax": DMAX,
        "p_max": MODEL_P_MAX,
        "support_hash": MODEL_SUPPORT_HASH,
        "support": list(SUPPORT_IDS),
        "Aend_definition": "last valid ILC column",
        "C2_definition": "literal second ILC column",
        "bandwidth": "5B",
        "B_MHz": 20,
        "Fs_MHz": 100,
        "state_count": STATE_COUNT,
        "allow_self": True,
        "Common_B_sha256": COMMON_B_SHA256,
        "raw_manifest_before": raw_before,
        "raw_manifest_after": raw_after,
        "raw_unchanged": raw_before == raw_after,
        **counts,
        "self_rate": counts["N_self"] / STATE_COUNT,
        "fallback_rate": counts["N_fallback"] / STATE_COUNT,
        "valid_rate": counts["N_valid"] / STATE_COUNT,
        "fail_rate": counts["N_fail"] / STATE_COUNT,
        **delta_stats,
        "finite_retrieved_real_B_CNMSE": finite_real_stats,
        "shareability_threshold_dB": REAL_B_THRESHOLD_DB,
        "shareability_operator": "<",
        "fingerprint_distance_shape": list(distance.shape),
        "fingerprint_distance_dtype": str(distance.dtype),
        "historical_regression": historical,
        "regression_oracle": EXPECTED_COUNTS,
        "figure_metadata": figure_metadata,
        "worker_count": WORKER_COUNT,
        "blas_threads_per_worker": BLAS_THREADS_PER_WORKER,
        "model_search_started": False,
        "nested_cv_started": False,
        "self_first_runner_started": False,
        "outputs": {
            "excel": str(RESULT_ROOT / "scenario_2_k18_aend_c2_full_lut_retrieval_5b.xlsx"),
            "csv": str(RESULT_ROOT / "01_state_results.csv"),
            "summary_json": str(RESULT_ROOT / "02_retrieval_summary.json"),
            "model_definition": str(RESULT_ROOT / "03_model_definition.json"),
            "distance_matrix": str(RESULT_ROOT / "04_fingerprint_distance_matrix.npy"),
            "figure_01": str(RESULT_ROOT / "figure_01_full_metrics_vs_state.png"),
            "figure_02": str(RESULT_ROOT / "figure_02_retrieved_real_B_CNMSE_vs_state.png"),
        },
    }
    _write_json(RESULT_ROOT / "02_retrieval_summary.json", summary)
    _write_excel(
        state_frame,
        summary,
        model_definition,
        RESULT_ROOT / "scenario_2_k18_aend_c2_full_lut_retrieval_5b.xlsx",
    )
    summary_lines = [
        f"Task: {TASK_NAME}",
        "Status: " + str(summary["status"]),
        f"Model: {MODEL_CANDIDATE}; K={MODEL_K}; lambda={RIDGE_LAMBDA:.12g}; dmax={DMAX}",
        f"LUT: Aend; Query: C2; bandwidth=5B; states={STATE_COUNT}; allow_self=true",
        f"Support hash: {MODEL_SUPPORT_HASH}",
        f"Common-B SHA256: {COMMON_B_SHA256}",
        "",
        f"N_self={counts['N_self']}",
        f"N_fallback={counts['N_fallback']}",
        f"N_valid={counts['N_valid']}",
        f"N_fail={counts['N_fail']}",
        f"self_rate={summary['self_rate']}",
        f"valid_rate={summary['valid_rate']}",
        f"mean_abs(State_Q-State_R)={delta_stats['mean_abs_state_delta']}",
        f"median_abs(State_Q-State_R)={delta_stats['median_abs_state_delta']}",
        f"max_abs(State_Q-State_R)={delta_stats['max_abs_state_delta']}",
        f"Real-B CNMSE < -40 dB count={counts['N_valid']}",
        f"Regression oracle 270/149/419/6 exact counts: {counts == EXPECTED_COUNTS}",
        f"Historical State_Q exact match: {historical.get('query_state_match')}",
        f"Historical distance max absolute difference: {historical.get('distance_matrix_max_abs_difference')}",
        "",
        f"Excel: {RESULT_ROOT / 'scenario_2_k18_aend_c2_full_lut_retrieval_5b.xlsx'}",
        f"Figure 01: {RESULT_ROOT / 'figure_01_full_metrics_vs_state.png'}",
        f"Figure 02: {RESULT_ROOT / 'figure_02_retrieved_real_B_CNMSE_vs_state.png'}",
        f"CSV: {RESULT_ROOT / '01_state_results.csv'}",
        f"Validation: {RESULT_ROOT / '08_validation_checks.json'}",
        f"raw_unchanged={summary['raw_unchanged']}",
        "No basis search, lambda search, Self-First rerun, Nested CV, or final validation was started.",
    ]
    (RESULT_ROOT / "07_final_result_summary.txt").write_text("\n".join(summary_lines) + "\n", encoding="utf-8")
    _write_checkpoint("completed", completed_model_state_ids=list(range(STATE_COUNT)), retrieval_rows=STATE_COUNT, **counts, raw_unchanged=bool(summary["raw_unchanged"]))
    _append_log(f"OUTPUTS WRITTEN status={summary['status']} raw_unchanged={summary['raw_unchanged']} historical_Q_match={historical.get('query_state_match')}")
    if summary["status"] != "SUCCESS":
        raise RuntimeError(f"fixed K18 regression mismatch; inspect {RESULT_ROOT / '02_retrieval_summary.json'}")
    _append_handoff(
        f"完成固定 K18 Aend/C2 Full-425 LUT 检索：N_self={counts['N_self']}，N_fallback={counts['N_fallback']}，N_valid={counts['N_valid']}，N_fail={counts['N_fail']}；历史 State_Q 一致={historical.get('query_state_match')}；raw_unchanged={summary['raw_unchanged']}。未启动模型搜索、Nested CV 或 Self-First runner。结果目录：{RESULT_ROOT}"
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--allow-existing", action="store_true")
    args = parser.parse_args()
    summary = run(allow_existing=args.allow_existing)
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=_json_default))


if __name__ == "__main__":
    main()
