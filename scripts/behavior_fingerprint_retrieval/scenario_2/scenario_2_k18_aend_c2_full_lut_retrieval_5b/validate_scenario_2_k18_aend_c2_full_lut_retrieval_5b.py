"""Independent validator for the fixed K18 Full-425 retrieval task."""

# ruff: noqa: E402,E501

from __future__ import annotations

import ast
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

PROJECT_ROOT = next(
    parent
    for parent in Path(__file__).resolve().parents
    if (parent / "environment.yml").is_file() and (parent / "scripts").is_dir()
)
SCRIPTS_ROOT = PROJECT_ROOT / "scripts"
if str(SCRIPTS_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_ROOT))

from core.shared.module_registry import MODULES  # noqa: E402

from behavior_fingerprint_retrieval.scenario_2.scenario_2_k18_aend_c2_full_lut_retrieval_5b.run_scenario_2_k18_aend_c2_full_lut_retrieval_5b import (  # noqa: E402
    COMMON_B_SHA256,
    EXPECTED_COUNTS,
    EXPECTED_RAW_BYTES,
    EXPECTED_RAW_FILE_COUNT,
    EXPECTED_RAW_MANIFEST,
    EXPECTED_RAW_MAT_COUNT,
    FROZEN_ROOT,
    HISTORICAL_DISTANCE_PATH,
    HISTORICAL_QUERY_PATH,
    MODEL_CANDIDATE,
    MODEL_K,
    MODEL_SUPPORT_HASH,
    RESULT_ROOT,
    RIDGE_LAMBDA,
    STATE_COUNT,
    SUPPORT_IDS,
    raw_manifest,
)

TASK_PACKAGE = "behavior_fingerprint_retrieval.scenario_2.scenario_2_k18_aend_c2_full_lut_retrieval_5b"
TASK_NAME = "scenario_2_k18_aend_c2_full_lut_retrieval_5b"


def _check(condition: bool, failures: list[str], name: str) -> bool:
    if not condition:
        failures.append(name)
    return condition


def _task_to_task_import_violations() -> list[str]:
    violations: list[str] = []
    task_root = Path(__file__).resolve().parent
    for path in task_root.glob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        for node in ast.walk(tree):
            module = ""
            if isinstance(node, ast.ImportFrom) and node.module:
                module = node.module
            elif isinstance(node, ast.Import):
                module = ",".join(alias.name for alias in node.names)
            if module.startswith("behavior_fingerprint_retrieval.scenario_") and not module.startswith(TASK_PACKAGE):
                violations.append(f"{path.name}: {module}")
            if "scenario_2_self_first_joint_basis_ridge_lut_retrieval_search_5b" in module:
                violations.append(f"{path.name}: forbidden Self-First task import")
    return violations


def validate() -> dict[str, object]:
    failures: list[str] = []
    summary_path = RESULT_ROOT / "02_retrieval_summary.json"
    model_path = RESULT_ROOT / "03_model_definition.json"
    csv_path = RESULT_ROOT / "01_state_results.csv"
    distance_path = RESULT_ROOT / "04_fingerprint_distance_matrix.npy"
    excel_path = RESULT_ROOT / "scenario_2_k18_aend_c2_full_lut_retrieval_5b.xlsx"
    checkpoint_path = RESULT_ROOT / "09_checkpoint.json"
    for path in (
        summary_path,
        model_path,
        csv_path,
        distance_path,
        excel_path,
        checkpoint_path,
        RESULT_ROOT / "figure_01_full_metrics_vs_state.png",
        RESULT_ROOT / "figure_02_retrieved_real_B_CNMSE_vs_state.png",
        RESULT_ROOT / "07_final_result_summary.txt",
    ):
        _check(path.is_file() and path.stat().st_size > 0, failures, f"exists:{path.name}")

    checks: dict[str, object] = {}
    if failures:
        checks["required_outputs_present"] = False
        return {"task": "scenario_2_k18_aend_c2_full_lut_retrieval_5b", "pass": False, "checks": checks, "failures": failures}

    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    model = json.loads(model_path.read_text(encoding="utf-8"))
    checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
    frame = pd.read_csv(csv_path)
    distance = np.asarray(np.load(distance_path, allow_pickle=False), dtype=np.float64)
    checks["state_count"] = _check(len(frame) == STATE_COUNT, failures, "state_count")
    checks["state_R_canonical"] = _check(
        np.array_equal(frame["state_R"].to_numpy(dtype=int), np.arange(STATE_COUNT)),
        failures,
        "state_R_canonical",
    )
    checks["distance_matrix_shape"] = _check(distance.shape == (STATE_COUNT, STATE_COUNT), failures, "distance_matrix_shape")
    checks["distance_matrix_no_nan_or_posinf"] = _check(
        not np.isnan(distance).any() and not np.isposinf(distance).any(),
        failures,
        "distance_matrix_no_nan_or_posinf",
    )
    checks["top1_in_range"] = _check(
        bool(np.all((frame["retrieved_state_Q"] >= 0) & (frame["retrieved_state_Q"] < STATE_COUNT))),
        failures,
        "top1_in_range",
    )
    checks["support_count"] = _check(len(model.get("support", [])) == MODEL_K, failures, "support_count")
    checks["support_order"] = _check(
        tuple(item.get("basis_name") for item in model.get("support", [])) == SUPPORT_IDS,
        failures,
        "support_order",
    )
    checks["support_hash"] = _check(model.get("support_hash") == MODEL_SUPPORT_HASH, failures, "support_hash")
    checks["candidate"] = _check(model.get("candidate") == MODEL_CANDIDATE, failures, "candidate")
    checks["K"] = _check(int(model.get("K", -1)) == MODEL_K, failures, "K")
    checks["lambda"] = _check(abs(float(model.get("lambda", np.nan)) - RIDGE_LAMBDA) <= 1e-24, failures, "lambda")
    checks["Common_B_SHA"] = _check(summary.get("Common_B_sha256") == COMMON_B_SHA256, failures, "Common_B_SHA")
    checks["checkpoint_completed"] = _check(checkpoint.get("phase") == "completed", failures, "checkpoint_completed")
    checks["checkpoint_counts"] = _check(
        all(int(checkpoint.get(key, -1)) == value for key, value in EXPECTED_COUNTS.items()),
        failures,
        "checkpoint_counts",
    )

    observed_counts = {
        "N_self": int((frame["retrieval_class"] == "SELF").sum()),
        "N_fallback": int((frame["retrieval_class"] == "FALLBACK").sum()),
        "N_valid": int(frame["retrieval_class"].isin(["SELF", "FALLBACK"]).sum()),
        "N_fail": int((frame["retrieval_class"] == "FAIL").sum()),
    }
    checks["regression_counts_270_149_419_6"] = _check(observed_counts == EXPECTED_COUNTS, failures, "regression_counts_270_149_419_6")
    checks["class_logic"] = _check(
        bool(
            np.all(frame.loc[frame["self_hit"], "retrieval_class"] == "SELF")
            and np.all(frame.loc[~frame["self_hit"] & frame["real_B_shareable_lt_minus40"], "retrieval_class"] == "FALLBACK")
            and np.all(frame.loc[~frame["self_hit"] & ~frame["real_B_shareable_lt_minus40"], "retrieval_class"] == "FAIL")
        ),
        failures,
        "class_logic",
    )
    required_finite_columns = [
        "nmse_withoutdpd_dB",
        "acpr_low_withoutdpd_dBc",
        "acpr_high_withoutdpd_dBc",
        "acpr_withoutdpd_mean_dBc",
        "Y_Aend_train_NMSE_dB",
        "Y_Aend_B_NMSE_dB",
        "Y_C2_train_NMSE_dB",
        "Y_C2_B_NMSE_dB",
        "retrieval_fingerprint_CNMSE_dB",
        "retrieved_state_abs_delta",
    ]
    checks["required_metrics_finite"] = _check(
        all(np.isfinite(frame[column].to_numpy(dtype=float)).all() for column in required_finite_columns),
        failures,
        "required_metrics_finite",
    )
    checks["real_B_no_nan"] = _check(
        not frame["retrieved_real_B_CNMSE_dB"].isna().any(),
        failures,
        "real_B_no_nan",
    )

    current_raw = raw_manifest()
    expected_raw = {
        "sha256": EXPECTED_RAW_MANIFEST,
        "file_count": EXPECTED_RAW_FILE_COUNT,
        "mat_count": EXPECTED_RAW_MAT_COUNT,
        "bytes": EXPECTED_RAW_BYTES,
    }
    checks["raw_manifest_expected"] = _check(current_raw == expected_raw, failures, "raw_manifest_expected")
    checks["raw_manifest_unchanged"] = _check(summary.get("raw_manifest_before") == current_raw == summary.get("raw_manifest_after"), failures, "raw_manifest_unchanged")

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(excel_path, read_only=False, data_only=False)
        checks["excel_sheets"] = _check(
            set(workbook.sheetnames) >= {"state_results", "summary", "model_definition"},
            failures,
            "excel_sheets",
        )
        checks["excel_row_count"] = _check(workbook["state_results"].max_row == STATE_COUNT + 1, failures, "excel_row_count")
        checks["excel_freeze"] = _check(workbook["state_results"].freeze_panes == "A2", failures, "excel_freeze")
        workbook.close()
    except Exception:
        checks["excel_readable"] = False
        failures.append("excel_readable")

    checks["task_to_task_imports"] = _check(not _task_to_task_import_violations(), failures, "task_to_task_imports")
    checks["no_search_started"] = _check(
        summary.get("model_search_started") is False
        and summary.get("nested_cv_started") is False
        and summary.get("self_first_runner_started") is False,
        failures,
        "no_search_started",
    )
    checks["module_registry_mirror"] = _check(
        all(
            sorted(path.name for path in (PROJECT_ROOT / root).iterdir() if path.is_dir()) == sorted(MODULES)
            for root in ("scripts", "results", "work_logs")
        ),
        failures,
        "module_registry_mirror",
    )
    checks["no_results_shared"] = _check(
        not any((PROJECT_ROOT / "results" / module / "shared").exists() for module in MODULES),
        failures,
        "no_results_shared",
    )
    checks["no_work_logs_shared"] = _check(
        not any((PROJECT_ROOT / "work_logs" / module / "shared").exists() for module in MODULES),
        failures,
        "no_work_logs_shared",
    )

    historical = summary.get("historical_regression", {})
    checks["historical_State_Q_match"] = _check(bool(historical.get("query_state_match")), failures, "historical_State_Q_match")
    checks["historical_class_match"] = _check(bool(historical.get("class_match")), failures, "historical_class_match")
    checks["historical_distance_match"] = _check(
        bool(historical.get("distance_matrix_match_at_1e-8")),
        failures,
        "historical_distance_match",
    )
    checks["figure_01"] = _check((RESULT_ROOT / "figure_01_full_metrics_vs_state.png").stat().st_size > 0, failures, "figure_01")
    checks["figure_02"] = _check((RESULT_ROOT / "figure_02_retrieved_real_B_CNMSE_vs_state.png").stat().st_size > 0, failures, "figure_02")

    report = {
        "task": TASK_NAME,
        "pass": not failures,
        "checks": checks,
        "failures": failures,
        "observed_counts": observed_counts,
        "raw_manifest": current_raw,
        "frozen_root": str(FROZEN_ROOT),
        "historical_query_path": str(HISTORICAL_QUERY_PATH),
        "historical_distance_path": str(HISTORICAL_DISTANCE_PATH),
    }
    (RESULT_ROOT / "08_validation_checks.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )
    log_path = PROJECT_ROOT / "work_logs" / "behavior_fingerprint_retrieval" / "scenario_2" /TASK_NAME / "execution_log.txt"
    with log_path.open("a", encoding="utf-8") as handle:
        handle.write(f"Validation: {'PASS' if not failures else 'FAIL'}; checks={len(checks)}; failures={failures}\n")
    if not failures:
        handoff = PROJECT_ROOT / "work_logs" / "core" / "codex_handoff" / "codex_handoff.txt"
        with handoff.open("a", encoding="utf-8") as handle:
            handle.write(
                f"\n2026-09-18 | 新任务：{TASK_NAME}\n"
                f"独立 validator PASS；固定 K18、lambda={RIDGE_LAMBDA:.12g}、Aend/C2、Full-425、Common-B、Real-B 严格 < -40 dB 及历史 State_Q/距离矩阵回归均通过；raw manifest unchanged；未启动模型搜索、Self-First runner 或 Nested CV。\n"
            )
    return report


def main() -> None:
    report = validate()
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    raise SystemExit(0 if report["pass"] else 1)


if __name__ == "__main__":
    main()
