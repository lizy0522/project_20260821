"""Run frozen Envelope18 Full-425 common-B LUT retrieval."""

from __future__ import annotations

import json
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

# Set BLAS limits before NumPy/SciPy imports in spawned workers.
# ruff: noqa: E402
for _thread_variable in (
    "OMP_NUM_THREADS",
    "OPENBLAS_NUM_THREADS",
    "MKL_NUM_THREADS",
    "NUMEXPR_NUM_THREADS",
    "VECLIB_MAXIMUM_THREADS",
):
    os.environ[_thread_variable] = "1"

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd

PROJECT_ROOT = next(
    parent
    for parent in Path(__file__).resolve().parents
    if (parent / "environment.yml").is_file() and (parent / "scripts").is_dir()
)
SCRIPTS_ROOT = PROJECT_ROOT / "scripts"
PROJECT_ROOT = SCRIPTS_ROOT.parent
if str(SCRIPTS_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_ROOT))

from behavior_modeling.shared.basis_function_selection.all425_envelope18_evaluation import (  # noqa: E402
    OLS_CONDITION_HARD_LIMIT,
)

from behavior_fingerprint_retrieval.shared.envelope18_commonB_full_lut import (  # noqa: E402
    BLAS_THREADS_PER_WORKER,
    COMMON_B_SOURCE_STATE,
    DMAX,
    FINAL_SUPPORT_IDS,
    HARD20_IDS,
    REAL_B_THRESHOLD_DB,
    RESULT_ROOT,
    STATE_COUNT,
    TASK_NAME,
    VALID_B_LENGTH,
    WORKER_COUNT,
    _finite_summary,
    _worker_init,
    build_common_b_probe,
    compute_cnmse_matrix,
    model_worker_entry,
    raw_manifest_gate,
    top1_retrieval,
    verify_runtime_contract,
)

WORK_LOG = (
    PROJECT_ROOT / "work_logs" / "behavior_fingerprint_retrieval" / "scenario_2" /TASK_NAME / "execution_log.txt"  # noqa: E501
)
HANDOFF_LOG = PROJECT_ROOT / "work_logs" / "core" / "codex_handoff" / "codex_handoff.txt"
CHECKPOINT_PATH = RESULT_ROOT / "11_checkpoint.json"


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%S%z")


def _append_log(path: Path, message: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(message)


def _write_checkpoint(
    *,
    phase: str,
    support_digest: str,
    common_b_digest: str,
    raw: dict[str, object],
    completed_model_state_ids: list[int],
    completed_retrieval_rows: int,
    model_frozen: bool,
    test_unlocked: bool,
) -> None:
    CHECKPOINT_PATH.write_text(
        json.dumps(
            {
                "phase": phase,
                "completed_model_state_ids": sorted(completed_model_state_ids),
                "completed_retrieval_rows": completed_retrieval_rows,
                "support_hash": support_digest,
                "K": len(FINAL_SUPPORT_IDS),
                "lambda": 0.0,
                "dmax": DMAX,
                "common_B_sha256": common_b_digest,
                "common_B_source_state": COMMON_B_SOURCE_STATE,
                "state_count": STATE_COUNT,
                "allow_self_retrieval": True,
                "train_bounds": [0, 16_384],
                "test_bounds": [16_384, 24_576],
                "abc_bounds": [0, 12_288, 17_203, 24_576],
                "worker_count": WORKER_COUNT,
                "blas_threads_per_worker": BLAS_THREADS_PER_WORKER,
                "model_frozen": model_frozen,
                "test_unlocked": test_unlocked,
                "raw_manifest": raw,
            },
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        ),
        encoding="utf-8",
    )


def _validate_resume_checkpoint(
    support_digest: str, common_b_digest: str, raw: dict[str, object]
) -> None:
    if not CHECKPOINT_PATH.is_file():
        return
    value = json.loads(CHECKPOINT_PATH.read_text(encoding="utf-8"))
    if value.get("support_hash") != support_digest:
        raise RuntimeError("Resume support hash differs from frozen Envelope18")
    if value.get("common_B_sha256") != common_b_digest:
        raise RuntimeError("Resume common-B hash differs from current probe")
    if value.get("raw_manifest") != raw:
        raise RuntimeError("Resume raw manifest differs from current raw tree")
    if int(value.get("K", -1)) != 18 or float(value.get("lambda", -1.0)) != 0.0:
        raise RuntimeError("Resume K/lambda differs from frozen model")
    if int(value.get("dmax", -1)) != DMAX:
        raise RuntimeError("Resume dmax differs from frozen model")
    if value.get("abc_bounds") != [0, 12_288, 17_203, 24_576]:
        raise RuntimeError("Resume ABC bounds changed")
    if value.get("allow_self_retrieval") is not True:
        raise RuntimeError("Resume self-retrieval policy changed")


def _write_common_b_metadata(path: Path, metadata: dict[str, object]) -> None:
    lines = [
        "Common-B probe metadata",
        *[f"{key}: {value}" for key, value in metadata.items()],
        "",
        "The probe is constructed once from canonical State 0 xin B and reused by all "
        "425 Aend/C2 models.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_excel(frame: pd.DataFrame, path: Path) -> None:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    excel_frame = frame.copy()
    if "retrieved_real_B_CNMSE_dB" in excel_frame:
        values = excel_frame["retrieved_real_B_CNMSE_dB"].to_numpy(dtype=float)
        excel_frame["retrieved_real_B_CNMSE_dB"] = [
            "-Inf" if np.isneginf(value) else float(value) for value in values
        ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All425"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{chr(64 + min(len(excel_frame.columns), 26))}{len(excel_frame) + 1}"
    )
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for column_index, column in enumerate(excel_frame.columns, start=1):
        cell = sheet.cell(row=1, column=column_index, value=column)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row_index, row in enumerate(excel_frame.itertuples(index=False, name=None), start=2):
        for column_index, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if isinstance(value, (float, np.floating)):
                cell.number_format = "0.000000"
    for column_index, column in enumerate(excel_frame.columns, start=1):
        width = max(len(str(column)) + 2, 12)
        if len(excel_frame):
            width = min(
                max(width, max(len(str(value)) for value in excel_frame[column].head(50)) + 2), 34
            )
        sheet.column_dimensions[chr(64 + column_index)].width = width
    workbook.save(path)
    loaded = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = loaded["All425"]
        if worksheet.max_row != STATE_COUNT + 1:
            raise RuntimeError(f"Excel All425 rows={worksheet.max_row}, expected {STATE_COUNT + 1}")
        if worksheet.max_column != len(excel_frame.columns):
            raise RuntimeError("Excel All425 column count differs from CSV")
    finally:
        loaded.close()


def _write_six_metric_plot(frame: pd.DataFrame, path: Path, *, annotate: bool) -> None:
    ordered = frame.sort_values("State_n_R")
    x = ordered["State_n_R"].to_numpy(dtype=int)
    specs = (
        ("nmse_withoutdpd_dB", "NMSE without DPD", "#7f7f7f", "-"),
        ("Y_Aend_train_NMSE_dB", "Y-Aend train", "#1f77b4", "-"),
        ("Y_Aend_B_NMSE_dB", "Y-Aend → B", "#2ca02c", "-"),
        ("Y_C2_train_NMSE_dB", "Y-C2 train", "#9467bd", "-"),
        ("Y_C2_B_NMSE_dB", "Y-C2 → B", "#8c564b", "-"),
        ("retrieved_real_B_CNMSE_dB", "Retrieved Real-B CNMSE", "#d62728", "o-"),
    )
    plot_values = []
    for column, _, _, _ in specs:
        values = ordered[column].to_numpy(dtype=float)
        if column == "retrieved_real_B_CNMSE_dB":
            values = np.where(np.isfinite(values), np.maximum(values, -55.0), -55.0)
        plot_values.append(values)
    lower = float(min(np.min(values) for values in plot_values) - 1.0)
    upper = float(max(np.max(values) for values in plot_values) + 1.0)
    fig, ax = plt.subplots(figsize=(36, 12), dpi=300)
    for (column, label, color, style), values in zip(specs, plot_values, strict=True):
        marker = "o" if style == "o-" else None
        ax.plot(x, values, color=color, linewidth=0.9, marker=marker, markersize=2.4, label=label)
        if annotate and column == "retrieved_real_B_CNMSE_dB":
            for index, (state_id, value, query_id) in enumerate(
                zip(x, values, ordered["State_n_Q"].to_numpy(dtype=int), strict=True)
            ):
                offset = 2.0 + float(index % 3) * 2.5
                ax.annotate(
                    f"Q{query_id:03d}",
                    (state_id, value),
                    xytext=(0, offset),
                    textcoords="offset points",
                    fontsize=4,
                    rotation=90,
                    ha="center",
                    va="bottom",
                    color=color,
                )
    ax.axhline(TARGET_DB, color="#111111", linestyle="--", linewidth=1.1, label="-40 dB Threshold")
    ax.set_xlim(0, STATE_COUNT - 1)
    ax.set_ylim(lower, upper)
    ax.set_xticks(
        [0, 25, 50, 75, 100, 125, 150, 175, 200, 225, 250, 275, 300, 325, 350, 375, 400, 424]
    )
    ax.set_xlabel("Load State")
    ax.set_ylabel("NMSE / CNMSE (dB)")
    ax.set_title("Envelope18 Full-LUT Retrieval Across 425 Load States")
    ax.grid(alpha=0.22)
    ax.legend(frameon=False, ncol=3, loc="best")
    if annotate:
        ax.text(
            0.01,
            0.01,
            "Values below -55 dB are clipped for visualization only.",
            transform=ax.transAxes,
            fontsize=8,
            va="bottom",
        )
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


TARGET_DB = -40.0


def _build_model_quality_frame(model_rows: pd.DataFrame) -> pd.DataFrame:
    return model_rows[
        [
            "state_id",
            "funMng",
            "funAng",
            "secMng",
            "secAng",
            "Vm",
            "Pin",
            "ilc_A_end",
            "C2_available",
            "Y_Aend_train_NMSE_dB",
            "Y_Aend_B_NMSE_dB",
            "Y_Aend_rank",
            "Y_Aend_condition_number",
            "Y_Aend_coefficient_norm",
            "Y_Aend_finite",
            "Y_C2_train_NMSE_dB",
            "Y_C2_B_NMSE_dB",
            "Y_C2_rank",
            "Y_C2_condition_number",
            "Y_C2_coefficient_norm",
            "Y_C2_finite",
        ]
    ].copy()


def _quality_summary(frame: pd.DataFrame) -> dict[str, object]:
    output: dict[str, object] = {}
    for column in (
        "Y_Aend_train_NMSE_dB",
        "Y_Aend_B_NMSE_dB",
        "Y_C2_train_NMSE_dB",
        "Y_C2_B_NMSE_dB",
    ):
        values = frame[column].to_numpy(dtype=float)
        output[column] = {
            "pass_count_lt_minus40": int(np.count_nonzero(values < TARGET_DB)),
            "median_dB": float(np.median(values)),
            "Q95_dB": float(np.quantile(values, 0.95)),
            "worst_dB": float(np.max(values)),
        }
    return output


def _retrieval_summary(
    frame: pd.DataFrame,
    fingerprint_distance: np.ndarray,
    real_b_values: np.ndarray,
) -> dict[str, object]:
    finite_real = _finite_summary(real_b_values)
    finite_fingerprint = _finite_summary(
        frame["retrieval_fingerprint_CNMSE_dB"].to_numpy(dtype=float)
    )
    abs_delta = frame["State_index_abs_delta"].to_numpy(dtype=float)
    return {
        "state_count": STATE_COUNT,
        "exact_state_hit_count": int(frame["is_exact_state_hit"].sum()),
        "exact_state_hit_rate": float(frame["is_exact_state_hit"].mean()),
        "retrieved_real_B_pass_count": int(frame["retrieved_real_B_pass"].sum()),
        "retrieved_real_B_pass_rate": float(frame["retrieved_real_B_pass"].mean()),
        "retrieval_fingerprint_CNMSE": finite_fingerprint,
        "retrieved_real_B_CNMSE_finite_only": finite_real,
        "retrieved_real_B_exact_self_count": int(np.count_nonzero(np.isneginf(real_b_values))),
        "State_index_abs_delta": {
            "median": float(np.median(abs_delta)),
            "mean": float(np.mean(abs_delta)),
            "Q90": float(np.quantile(abs_delta, 0.90)),
            "max": float(np.max(abs_delta)),
        },
        "distance_shape": list(fingerprint_distance.shape),
    }


def _summary_text(payload: dict[str, object]) -> str:
    quality = payload["model_quality_summary"]
    retrieval = payload["retrieval_summary"]
    fingerprint_stats = json.dumps(
        retrieval["retrieval_fingerprint_CNMSE"], ensure_ascii=False, sort_keys=True
    )
    real_b_stats = json.dumps(
        retrieval["retrieved_real_B_CNMSE_finite_only"], ensure_ascii=False, sort_keys=True
    )
    delta_stats = json.dumps(retrieval["State_index_abs_delta"], ensure_ascii=False, sort_keys=True)
    lines = [
        f"Task: {TASK_NAME}",
        "Frozen model: Envelope18; K=18; lambda=0; dmax=2; solver=OLS.",
        "LUT mode: Full 425-state LUT; no clustering; self retrieval allowed.",
        "LUT fingerprint: Y-Aend -> common B.",
        "Online query: Y-C2 -> common B.",
        f"Support hash: {payload['support_hash']}",
        f"Common-B source state: {COMMON_B_SOURCE_STATE}",
        f"Common-B SHA256: {payload['common_b_metadata']['sha256']}",
        "",
        "Core retrieval result:",
        f"Exact state hits: {retrieval['exact_state_hit_count']}/{STATE_COUNT}",
        "Retrieved Real-B CNMSE < -40 dB: "
        f"{retrieval['retrieved_real_B_pass_count']}/{STATE_COUNT}",
        f"Fingerprint CNMSE finite-only: {fingerprint_stats}",
        f"Retrieved Real-B CNMSE finite-only: {real_b_stats}",
        f"Exact self-hit -Inf count: {retrieval['retrieved_real_B_exact_self_count']}",
        f"State index abs delta: {delta_stats}",
        "State index delta is only an auxiliary canonical-order difference, not a physical "
        "load distance.",
        "",
        "Model quality:",
    ]
    for key, value in quality.items():
        lines.append(f"{key}: {json.dumps(value, ensure_ascii=False, sort_keys=True)}")
    lines.extend(
        [
            "",
            "Strict Real-B success uses retrieved_real_B_CNMSE_dB < -40 dB; Q!=R is still "
            "a valid behavior match.",
            "Model coefficients are state-specific, but the Envelope18 structure, lambda, "
            "and common-B probe are unified.",
            "Raw no-DPD NMSE/ACPR values are copied from the formal MAT fields; ACPR is not "
            "plotted in the six-metric figure.",
            "No Aend/C2 column scan, multi-ILC fusion, model retuning, Top-k fallback, "
            "clustering, Type-III compression, DPD replay, or low-bandwidth task was run.",
            "Raw snapshot before: "
            f"{json.dumps(payload['raw_before'], ensure_ascii=False, sort_keys=True)}",
            "Raw snapshot after: "
            f"{json.dumps(payload['raw_after'], ensure_ascii=False, sort_keys=True)}",
            f"raw_data_modified: {payload['raw_before'] != payload['raw_after']}",
        ]
    )
    return "\n".join(lines) + "\n"


def _run(*, resume: bool = False) -> dict[str, object]:
    existing_output = RESULT_ROOT.exists() and any(RESULT_ROOT.iterdir())
    if existing_output and not resume:
        raise RuntimeError(
            "Result directory is not empty; use --resume only for the same incomplete task: "
            f"{RESULT_ROOT}"
        )
    RESULT_ROOT.mkdir(parents=True, exist_ok=True)
    WORK_LOG.parent.mkdir(parents=True, exist_ok=True)
    raw_before = raw_manifest_gate()
    terms, support_indices, support_digest = verify_runtime_contract()
    common_b, common_phi, common_meta = build_common_b_probe()
    _validate_resume_checkpoint(support_digest, common_meta["sha256"], raw_before)
    (RESULT_ROOT / "00_task_definition.txt").write_text(
        "\n".join(
            [
                f"Task: {TASK_NAME}",
                "Frozen Envelope18 Aend/C2 common-B Full-425 LUT retrieval.",
                "Data chains: Aend/C2 use ILC pairs; Real-B uses xin -> yout_withoutdpd_ori.",
                "LUT: 425 Aend fingerprints; Query: 425 C2 fingerprints; self retrieval allowed.",
                f"Support IDs: {list(FINAL_SUPPORT_IDS)}",
                f"Support hash: {support_digest}",
                "Preprocessing: canonical full-record ABC alignment/gain; no pairwise "
                "re-alignment.",
                "Selection: none; lambda/dmax/support are frozen.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    _write_common_b_metadata(RESULT_ROOT / "01_common_B_probe_metadata.txt", common_meta)
    _append_log(
        WORK_LOG,
        "\n"
        f"[{_now()}] Start {TASK_NAME}\n"
        "Mode: frozen Envelope18 Full-425 LUT retrieval; no model selection.\n"
        f"Raw snapshot before: {json.dumps(raw_before, ensure_ascii=False, sort_keys=True)}\n"
        f"Support hash: {support_digest}; common-B sha256={common_meta['sha256']}; "
        f"workers={WORKER_COUNT}; BLAS threads/worker=1.\n",
    )
    _write_checkpoint(
        phase="setup_complete",
        support_digest=support_digest,
        common_b_digest=str(common_meta["sha256"]),
        raw=raw_before,
        completed_model_state_ids=[],
        completed_retrieval_rows=0,
        model_frozen=True,
        test_unlocked=False,
    )

    model_results: dict[int, dict[str, object]] = {}
    context = __import__("multiprocessing").get_context("spawn")
    with ProcessPoolExecutor(
        max_workers=WORKER_COUNT,
        mp_context=context,
        initializer=_worker_init,
        initargs=(tuple(support_indices), np.asarray(common_phi, dtype=np.complex128)),
    ) as executor:
        futures = {
            executor.submit(model_worker_entry, state_id): state_id
            for state_id in range(STATE_COUNT)
        }
        for completed, future in enumerate(as_completed(futures), start=1):
            state_id = futures[future]
            model_results[state_id] = future.result()
            if completed % 50 == 0 or completed == STATE_COUNT:
                print(f"[MODEL] {completed}/{STATE_COUNT}", flush=True)
                _write_checkpoint(
                    phase="model_progress",
                    support_digest=support_digest,
                    common_b_digest=str(common_meta["sha256"]),
                    raw=raw_before,
                    completed_model_state_ids=sorted(model_results),
                    completed_retrieval_rows=0,
                    model_frozen=True,
                    test_unlocked=False,
                )
    if sorted(model_results) != list(range(STATE_COUNT)):
        raise RuntimeError("Aend/C2 model preparation did not cover all 425 states")
    model_rows = []
    a_theta = np.empty((STATE_COUNT, len(support_indices)), dtype=np.complex128)
    c_theta = np.empty_like(a_theta)
    real_b = np.empty((STATE_COUNT, VALID_B_LENGTH), dtype=np.complex128)
    for state_id in range(STATE_COUNT):
        result = model_results[state_id]
        model_rows.append(
            {
                key: value
                for key, value in result.items()
                if key not in {"Aend_theta", "C2_theta", "real_B_valid"}
            }
        )
        a_theta[state_id] = result["Aend_theta"]
        c_theta[state_id] = result["C2_theta"]
        real_b[state_id] = result["real_B_valid"]
    model_frame = pd.DataFrame(model_rows).sort_values("state_id").reset_index(drop=True)
    if len(model_frame) != STATE_COUNT or not np.array_equal(
        model_frame["state_id"].to_numpy(dtype=int), np.arange(STATE_COUNT)
    ):
        raise RuntimeError("Model metrics are not canonical 0...424")
    if not np.all(np.isfinite(a_theta)) or not np.all(np.isfinite(c_theta)):
        raise RuntimeError("Aend/C2 theta contains NaN or Inf")
    if not np.all(np.isfinite(real_b)):
        raise RuntimeError("Real-B library contains NaN or Inf")
    if not np.all(model_frame["C2_available"].astype(bool)):
        raise RuntimeError("C2 is not available for all 425 states")
    if not np.all(model_frame["Y_Aend_rank"].to_numpy(dtype=int) == 18):
        raise RuntimeError("Aend rank gate failed")
    if not np.all(model_frame["Y_C2_rank"].to_numpy(dtype=int) == 18):
        raise RuntimeError("C2 rank gate failed")
    if not np.all(model_frame[["Y_Aend_finite", "Y_C2_finite"]].astype(bool).to_numpy()):
        raise RuntimeError("Aend/C2 finite gate failed")
    conditions = model_frame[["Y_Aend_condition_number", "Y_C2_condition_number"]].to_numpy(
        dtype=float
    )
    if not np.all(np.isfinite(conditions)) or np.max(conditions) > OLS_CONDITION_HARD_LIMIT:
        raise RuntimeError("Aend/C2 condition-number gate failed")
    quality_frame = _build_model_quality_frame(model_frame)
    quality_frame.to_csv(RESULT_ROOT / "02_model_quality_all425.csv", index=False)
    _write_checkpoint(
        phase="model_fit_complete",
        support_digest=support_digest,
        common_b_digest=str(common_meta["sha256"]),
        raw=raw_before,
        completed_model_state_ids=list(range(STATE_COUNT)),
        completed_retrieval_rows=0,
        model_frozen=True,
        test_unlocked=False,
    )

    lut_fingerprints = np.asarray(common_phi @ a_theta.T, dtype=np.complex128).T
    query_fingerprints = np.asarray(common_phi @ c_theta.T, dtype=np.complex128).T
    if lut_fingerprints.shape != (STATE_COUNT, VALID_B_LENGTH):
        raise RuntimeError(f"LUT fingerprint shape failed: {lut_fingerprints.shape}")
    if query_fingerprints.shape != (STATE_COUNT, VALID_B_LENGTH):
        raise RuntimeError(f"Query fingerprint shape failed: {query_fingerprints.shape}")
    if not np.all(np.isfinite(lut_fingerprints)) or not np.all(np.isfinite(query_fingerprints)):
        raise RuntimeError("LUT/Query fingerprints contain NaN or Inf")
    fingerprint_distance = compute_cnmse_matrix(query_fingerprints, lut_fingerprints)
    if fingerprint_distance.shape != (STATE_COUNT, STATE_COUNT):
        raise RuntimeError("Fingerprint CNMSE matrix shape failed")
    np.save(RESULT_ROOT / "05_fingerprint_cnmse_matrix.npy", fingerprint_distance)
    selected, selected_distance, true_rank, tie_count = top1_retrieval(fingerprint_distance)
    real_b_distance = compute_cnmse_matrix(real_b, real_b)
    retrieved_real = real_b_distance[np.arange(STATE_COUNT), selected]
    retrieval_frame = model_frame.copy()
    retrieval_frame["State_n_R"] = retrieval_frame["state_id"].astype(int)
    retrieval_frame["State_n_Q"] = selected.astype(int)
    retrieval_frame["retrieval_fingerprint_CNMSE_dB"] = selected_distance
    retrieval_frame["retrieved_real_B_CNMSE_dB"] = retrieved_real
    retrieval_frame["State_index_delta"] = selected.astype(int) - retrieval_frame["State_n_R"]
    retrieval_frame["State_index_abs_delta"] = np.abs(retrieval_frame["State_index_delta"])
    retrieval_frame["is_exact_state_hit"] = selected == np.arange(STATE_COUNT)
    retrieval_frame["retrieved_real_B_pass"] = retrieved_real < REAL_B_THRESHOLD_DB
    retrieval_frame["fingerprint_true_state_rank"] = true_rank
    retrieval_frame["fingerprint_tie_count"] = tie_count
    retrieval_frame["is_hard20"] = retrieval_frame["state_id"].isin(HARD20_IDS)
    retrieval_frame = retrieval_frame.sort_values("State_n_R").reset_index(drop=True)
    if not np.array_equal(retrieval_frame["State_n_R"].to_numpy(dtype=int), np.arange(STATE_COUNT)):
        raise RuntimeError("Retrieval results are not ordered by State_n_R")
    model_columns = [
        "State_n_R",
        "funMng",
        "funAng",
        "secMng",
        "secAng",
        "nmse_withoutdpd_dB",
        "ACPR_withoutdpd_mean_dBc",
        "Y_Aend_train_NMSE_dB",
        "Y_Aend_B_NMSE_dB",
        "Y_C2_train_NMSE_dB",
        "Y_C2_B_NMSE_dB",
        "State_n_Q",
        "retrieved_real_B_CNMSE_dB",
        "retrieval_fingerprint_CNMSE_dB",
        "State_index_delta",
        "State_index_abs_delta",
        "is_exact_state_hit",
        "retrieved_real_B_pass",
    ]
    core_frame = retrieval_frame[model_columns].copy()
    retrieval_frame.to_csv(RESULT_ROOT / "03_all425_retrieval_metrics.csv", index=False)
    failed = core_frame.loc[~core_frame["retrieved_real_B_pass"].astype(bool)].copy()
    failed["failure_type"] = "retrieved_real_B_CNMSE_ge_minus40dB"
    failed.to_csv(RESULT_ROOT / "04_failed_real_B_states.csv", index=False)
    _write_excel(core_frame, RESULT_ROOT / "06_envelope18_full_lut_retrieval.xlsx")
    core_frame.to_csv(RESULT_ROOT / "07_envelope18_full_lut_retrieval.csv", index=False)
    _write_six_metric_plot(
        core_frame, RESULT_ROOT / "08_envelope18_full_lut_retrieval_six_metrics.png", annotate=True
    )
    _write_six_metric_plot(
        core_frame, RESULT_ROOT / "09_envelope18_full_lut_retrieval_clean.png", annotate=False
    )
    quality_summary = _quality_summary(model_frame)
    retrieval_summary = _retrieval_summary(core_frame, fingerprint_distance, retrieved_real)
    raw_after = raw_manifest_gate()
    payload = {
        "support_hash": support_digest,
        "common_b_metadata": common_meta,
        "model_quality_summary": quality_summary,
        "retrieval_summary": retrieval_summary,
        "raw_before": raw_before,
        "raw_after": raw_after,
    }
    (RESULT_ROOT / "10_final_result_summary.txt").write_text(
        _summary_text(payload), encoding="utf-8"
    )
    _write_checkpoint(
        phase="completed",
        support_digest=support_digest,
        common_b_digest=str(common_meta["sha256"]),
        raw=raw_after,
        completed_model_state_ids=list(range(STATE_COUNT)),
        completed_retrieval_rows=STATE_COUNT,
        model_frozen=True,
        test_unlocked=True,
    )
    return {
        "task": TASK_NAME,
        "exact_state_hit_count": int(core_frame["is_exact_state_hit"].sum()),
        "retrieved_real_B_pass_count": int(core_frame["retrieved_real_B_pass"].sum()),
        "failed_count": int(len(failed)),
        "raw_data_modified": raw_before != raw_after,
        "result_root": str(RESULT_ROOT),
        "support_hash": support_digest,
        "common_B_sha256": common_meta["sha256"],
    }


def main() -> None:
    try:
        result = _run(resume="--resume" in sys.argv[1:])
    except Exception as exc:
        _append_log(
            WORK_LOG,
            f"[{_now()}] FAILED {TASK_NAME}: {type(exc).__name__}: {exc}\n",
        )
        raise
    _append_log(
        WORK_LOG,
        f"[{_now()}] Completed {TASK_NAME}\n"
        f"Result: {json.dumps(result, ensure_ascii=False, sort_keys=True)}\n"
        "No re-selection, Ridge scan, dmax/p-order expansion, clustering, Type-III, "
        "DPD, or low-bandwidth task was run.\n",
    )
    _append_log(
        HANDOFF_LOG,
        f"\n{_now()} | 新任务：{TASK_NAME}\n"
        "完成冻结 Envelope18 + Aend/C2 + common-B + Full-425 LUT retrieval；"
        f"结果目录：{RESULT_ROOT}\n"
        f"结果摘要：{json.dumps(result, ensure_ascii=False, sort_keys=True)}\n"
        "使用 10 个 spawn worker、每 worker 1 个 BLAS 线程；Full425 LUT 不聚类、"
        "允许 self retrieval；"
        "未运行 support/lambda/dmax 调整、Type-III、DPD 或 low-bandwidth。\n",
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    __import__("multiprocessing").freeze_support()
    main()
