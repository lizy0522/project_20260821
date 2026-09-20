"""Evaluation-only Full425 common-B LUT retrieval for frozen Envelope19."""

from __future__ import annotations

import json
import os
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from multiprocessing import get_context
from pathlib import Path

# ruff: noqa: E402,E501

for _thread_variable in (
    "OMP_NUM_THREADS",
    "OPENBLAS_NUM_THREADS",
    "MKL_NUM_THREADS",
    "NUMEXPR_NUM_THREADS",
    "VECLIB_MAXIMUM_THREADS",
):
    os.environ[_thread_variable] = "1"

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MODULE_ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = next(
    parent
    for parent in Path(__file__).resolve().parents
    if (parent / "environment.yml").is_file() and (parent / "scripts").is_dir()
)
SCRIPTS_ROOT = PROJECT_ROOT / "scripts"
PROJECT_ROOT = SCRIPTS_ROOT.parent
if str(SCRIPTS_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_ROOT))

from behavior_modeling.shared.basis_function_selection.envelope_dictionary import (
    build_envelope_bank,  # noqa: E402
)
from data_management.shared import build_state_table  # noqa: E402

from behavior_fingerprint_retrieval.shared.envelope18_commonB_full_lut import (  # noqa: E402
    top1_retrieval,
)
from behavior_fingerprint_retrieval.shared.envelope19_c2endshared_commonB_full_lut import (  # noqa: E402
    EXPECTED_COMMON_B_SHA,
    EXPECTED_SUPPORT_HASH,
    MODEL_NAME,
    RESULT_ROOT,
    STATE_COUNT,
    TASK_NAME,
    VALID_B_LENGTH,
    _finite_summary,
    _worker_init,
    compute_cnmse_matrix,
    load_frozen_support,
    model_worker_entry,
    raw_manifest_gate,
    verify_common_b_contract,
)

E23_ROOT = PROJECT_ROOT / "results" / "behavior_fingerprint_retrieval" / "scenario_2" /"scenario_2_envelope23_ilcend_commonB_full_lut_retrieval_5B"
E23_METRICS = E23_ROOT / "03_envelope23_all425_retrieval_metrics.csv"
E18_ROOT = PROJECT_ROOT / "results" / "behavior_fingerprint_retrieval" / "scenario_2" /"scenario_2_envelope18_commonB_full_lut_retrieval_5B"
E18_METRICS = E18_ROOT / "03_all425_retrieval_metrics.csv"
CHECKPOINT_PATH = RESULT_ROOT / "13_checkpoint.json"
WORK_LOG = PROJECT_ROOT / "work_logs" / "behavior_fingerprint_retrieval" / "scenario_2" /TASK_NAME / "execution_log.txt"
HANDOFF_LOG = PROJECT_ROOT / "work_logs" / "core" / "codex_handoff" / "codex_handoff.txt"
THRESHOLD_DB = -40.0
MODEL_DEF_PATH = PROJECT_ROOT / "results" / "behavior_modeling" / "scenario_2" /"scenario_2_hard20_c2_ilcend_shared_envelope75_basis_selection_5B" / "07_final_model_definition.csv"


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%S%z")


def _append_log(path: Path, message: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(message)


def _checkpoint(**payload: object) -> None:
    CHECKPOINT_PATH.write_text(
        json.dumps(
            {
                "task_name": TASK_NAME,
                "mode": "evaluation_only",
                "model_name": MODEL_NAME,
                "K": 19,
                "dmax": 2,
                "lambda": 0.0,
                "solver": "OLS",
                "state_count": STATE_COUNT,
                "abc_bounds": [0, 12_288, 17_203, 24_576],
                "allow_self_retrieval": True,
                "common_B_source_state": 0,
                "common_B_sha256": EXPECTED_COMMON_B_SHA,
                "support_hash": EXPECTED_SUPPORT_HASH,
                "worker_count": 10,
                "blas_threads_per_worker": 1,
                "basis_selection_performed": False,
                "model_selection_performed": False,
                "ridge_scan_performed": False,
                "dmax_scan_performed": False,
                "lut_retrieval_performed": True,
                **payload,
            },
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        ),
        encoding="utf-8",
    )


def _write_excel(frame: pd.DataFrame, path: Path) -> None:
    headers = [
        "状态序号",
        "负载配置",
        "NMSE(NMSE_WITHOUTDPD)",
        "上下边带ACPR平均值withoutdpd",
        "A段建模精度",
        "AB段泛化精度",
        "C段建模精度",
        "CB段泛化精度",
        "命中的序号",
        "命中的CNMSE",
    ]

    def mismatch(value: object) -> str:
        number = int(round(float(value)))
        return "0" if number == 0 else f"{number / 100:.2f}".rstrip("0").rstrip(".")

    state_table = {int(row["state_id"]): row for row in build_state_table()}
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "retrieval_results"
    worksheet.append(headers)
    for _, row in frame.sort_values("State_n_R").iterrows():
        state_id = int(row["State_n_R"])
        state = state_table[state_id]
        load_config = (
            f"funMng={mismatch(state['funMng'])}, funAng={int(state['funAng'])}°, "
            f"secMng={mismatch(state['secMng'])}, secAng={int(state['secAng'])}°"
        )
        real_b = row["retrieved_real_B_CNMSE_dB"]
        hit_cnmse: object = "-Inf" if np.isneginf(float(real_b)) else float(real_b)
        worksheet.append(
            [
                state_id,
                load_config,
                float(row["nmse_withoutdpd_dB"]),
                float(row["ACPR_withoutdpd_mean_dBc"]),
                float(row["Y_Aend_train_NMSE_dB"]),
                float(row["Y_Aend_B_NMSE_dB"]),
                float(row["Y_C2_train_NMSE_dB"]),
                float(row["Y_C2_B_NMSE_dB"]),
                int(row["State_n_Q"]),
                hit_cnmse,
            ]
        )
    if worksheet.max_row != 426 or worksheet.max_column != 10:
        raise RuntimeError("Excel output shape is not 426x10")
    if [worksheet.cell(1, column).value for column in range(1, 11)] != headers:
        raise RuntimeError("Excel header contract failed")
    if [worksheet.cell(row, 1).value for row in range(2, 427)] != list(range(425)):
        raise RuntimeError("Excel State_R ordering is not 0...424")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2, max_row=426, min_col=1, max_col=10):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F1F1F")
            cell.alignment = Alignment(vertical="center")
        for cell in row[2:8]:
            cell.number_format = "0.000000000"
        row[9].number_format = "0.000000000"
    widths = [12, 46, 22, 28, 20, 20, 20, 20, 14, 20]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:J426"
    worksheet.sheet_view.showGridLines = False
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)

    reloaded = load_workbook(path, read_only=True, data_only=True)
    sheet = reloaded["retrieval_results"]
    if sheet.max_row != 426 or sheet.max_column != 10:
        raise RuntimeError("reloaded Excel shape is not 426x10")
    if [sheet.cell(1, column).value for column in range(1, 11)] != headers:
        raise RuntimeError("reloaded Excel headers changed")
    reloaded.close()


def _write_plot(
    frame: pd.DataFrame,
    path: Path,
    *,
    title: str = "Frozen Envelope19-C2EndShared Full425 common-B LUT retrieval",
) -> None:
    ordered = frame.sort_values("State_n_R")
    x = ordered["State_n_R"].to_numpy(dtype=int)
    specs = (
        ("nmse_withoutdpd_dB", "No-DPD NMSE", "#7f7f7f"),
        ("Y_Aend_train_NMSE_dB", "Aend Train", "#1f77b4"),
        ("Y_Aend_B_NMSE_dB", "Aend -> B", "#2ca02c"),
        ("Y_C2_train_NMSE_dB", "C2 Train", "#9467bd"),
        ("Y_C2_B_NMSE_dB", "C2 -> B", "#8c564b"),
        ("retrieved_real_B_CNMSE_dB", "Retrieved Real-B", "#d62728"),
    )
    finite_metric_values = []
    for column, _, _ in specs:
        values = ordered[column].to_numpy(dtype=float)
        finite_metric_values.extend(values[np.isfinite(values)].tolist())
    floor = min(-55.0, float(np.min(finite_metric_values)) - 1.0)
    ceiling = max(-40.0, float(np.max(finite_metric_values)) + 1.0)
    fig, ax = plt.subplots(figsize=(24, 8), dpi=300)
    for column, label, color in specs:
        values = ordered[column].to_numpy(dtype=float)
        plot_values = np.where(np.isfinite(values), values, floor)
        marker = "o" if column == "retrieved_real_B_CNMSE_dB" else None
        ax.plot(x, plot_values, linewidth=0.8, marker=marker, markersize=2.0, label=label, color=color)
        if column == "retrieved_real_B_CNMSE_dB":
            for state_id, query_value, display_value in zip(
                x,
                plot_values,
                ordered["State_n_Q"].to_numpy(dtype=int),
                strict=True,
            ):
                offset = 5 if state_id % 2 == 0 else -8
                ax.annotate(
                    str(int(display_value)),
                    (int(state_id), float(query_value)),
                    xytext=(0, offset),
                    textcoords="offset points",
                    fontsize=3.6,
                    rotation=90,
                    ha="center",
                    va="bottom" if offset > 0 else "top",
                    color=color,
                )
    ax.axhline(THRESHOLD_DB, color="#111111", linestyle="--", linewidth=1.0, label="-40 dB threshold")
    ax.set_xlim(0, STATE_COUNT - 1)
    ax.set_ylim(floor, ceiling)
    ax.set_xticks([0, 25, 50, 75, 100, 125, 150, 175, 200, 225, 250, 275, 300, 325, 350, 375, 400, 424])
    ax.set_xlabel("State ID")
    ax.set_ylabel("NMSE / CNMSE (dB)")
    ax.set_title(title)
    ax.grid(alpha=0.22)
    ax.legend(frameon=False, ncol=3, loc="upper left")
    ax.text(
        0.99,
        0.01,
        "-Inf exact self-hit values are placed at the plotting floor for display only; labels are State_Q.",
        transform=ax.transAxes,
        fontsize=7,
        ha="right",
        va="bottom",
    )
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


def _format_quality(frame: pd.DataFrame, model_name: str, prefix: str) -> list[dict[str, object]]:
    rows = []
    for metric in ("Aend_train", "Aend_B", "C2_train", "C2_B"):
        values = frame[f"Y_{metric}_NMSE_dB"].to_numpy(dtype=float)
        rows.append(
            {
                "model": model_name,
                "metric": metric,
                "pass_count_lt_minus40": int(np.count_nonzero(values < THRESHOLD_DB)),
                "median_dB": float(np.median(values)),
                "mean_dB": float(np.mean(values)),
                "Q95_dB": float(np.quantile(values, 0.95)),
                "worst_dB": float(np.max(values)),
            }
        )
    return rows


def _run(*, resume: bool = False) -> dict[str, object]:
    existing = RESULT_ROOT.exists() and any(RESULT_ROOT.iterdir())
    if existing and not resume:
        raise RuntimeError("result directory is not empty; use --resume")
    if resume:
        if not CHECKPOINT_PATH.is_file():
            # A setup-only failed attempt has no resumable state; rerun setup
            # in the same task directory and overwrite only this task's files.
            existing = False
    RESULT_ROOT.mkdir(parents=True, exist_ok=True)
    raw_before = raw_manifest_gate()
    terms, support, support_ids, support_digest, model_frame = load_frozen_support()
    common_b, common_meta = verify_common_b_contract()
    common_phi_full = build_envelope_bank(common_b, terms)
    common_phi = np.asarray(common_phi_full[:, support], dtype=np.complex128)
    if common_phi.shape != (VALID_B_LENGTH, 19):
        raise RuntimeError(f"Envelope19 common-B Phi shape is {common_phi.shape}")

    (RESULT_ROOT / "00_task_definition.txt").write_text(
        "\n".join(
            [
                f"Task: {TASK_NAME}",
                "Mode: evaluation_only.",
                f"Frozen model: {MODEL_NAME}; K=19; dmax=2; lambda=0; solver=OLS.",
                f"Support hash: {support_digest}",
                "Only changed variable versus the E23 Full-LUT protocol: E23 support -> Envelope19-C2EndShared.",
                "Aend/C2 semantics, ABC bounds, common-B, 425x425 CNMSE, Top-1 allow-self rule, and Real-B threshold are frozen.",
                "No support selection, Ridge, dmax scan, clustering, Type-III, DPD replay, or low-bandwidth task.",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    contract = {
        "task_name": TASK_NAME,
        "evaluation_only": True,
        "model_name": MODEL_NAME,
        "K": 19,
        "dmax": 2,
        "lambda": 0.0,
        "solver": "OLS",
        "support_hash": support_digest,
        "basis_ids": list(support_ids),
        "reference_e23_task": "scenario_2_envelope23_ilcend_commonB_full_lut_retrieval_5B",
        "abc_bounds": [0, 12_288, 17_203, 24_576],
        "allow_self": True,
        "state_count": STATE_COUNT,
        "common_B_expected_hash": EXPECTED_COMMON_B_SHA,
        "basis_selection": False,
        "ridge_scan": False,
        "dmax_scan": False,
        "lut_retrieval": True,
    }
    (RESULT_ROOT / "01_frozen_model_and_protocol_contract.json").write_text(
        json.dumps(contract, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (RESULT_ROOT / "02_commonB_metadata.json").write_text(
        json.dumps(
            {
                **{key: value for key, value in common_meta.items() if not isinstance(value, np.ndarray)},
                "same_as_E23_commonB": True,
                "expected_sha256": EXPECTED_COMMON_B_SHA,
                "Envelope19_support_hash": support_digest,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    _checkpoint(
        phase="setup_complete",
        support_hash=support_digest,
        common_B_hash=EXPECTED_COMMON_B_SHA,
        raw_manifest_before=raw_before,
        raw_manifest_after=None,
        completed_model_state_ids=[],
        completed_retrieval_rows=0,
        coefficients_completed=False,
        fingerprints_completed=False,
        retrieval_completed=False,
        excel_completed=False,
        plot_completed=False,
    )
    _append_log(
        WORK_LOG,
        "\n"
        f"[{_now()}] Start {TASK_NAME}\n"
        f"Frozen support hash={support_digest}; common-B sha={EXPECTED_COMMON_B_SHA}; raw={json.dumps(raw_before, sort_keys=True)}\n"
        "Protocol frozen from E23 Full-LUT retrieval; only support changed to Envelope19-C2EndShared.\n",
    )

    context = get_context("spawn")
    results: list[dict[str, object]] = []
    with ProcessPoolExecutor(
        max_workers=10,
        mp_context=context,
        initializer=_worker_init,
        initargs=(tuple(support),),
    ) as executor:
        futures = {executor.submit(model_worker_entry, state_id): state_id for state_id in range(STATE_COUNT)}
        for completed, future in enumerate(as_completed(futures), start=1):
            results.append(future.result())
            if completed % 50 == 0 or completed == STATE_COUNT:
                _checkpoint(
                    phase="model_progress",
                    support_hash=support_digest,
                    common_B_hash=EXPECTED_COMMON_B_SHA,
                    raw_manifest_before=raw_before,
                    raw_manifest_after=None,
                    completed_model_state_ids=sorted(int(item["State_n_R"]) for item in results),
                    completed_retrieval_rows=0,
                    coefficients_completed=False,
                    fingerprints_completed=False,
                    retrieval_completed=False,
                    excel_completed=False,
                    plot_completed=False,
                )
                print(f"[E19 MODEL] {completed}/{STATE_COUNT}", flush=True)
    results.sort(key=lambda item: int(item["State_n_R"]))
    model_frame = pd.DataFrame(
        [
            {key: value for key, value in item.items() if key not in {"theta_Aend", "theta_C2", "real_B_valid"}}
            for item in results
        ]
    )
    model_frame.to_csv(RESULT_ROOT / "03_all425_model_quality.csv", index=False)
    theta_a = np.stack([item["theta_Aend"] for item in results], axis=0).astype(np.complex128)
    theta_c = np.stack([item["theta_C2"] for item in results], axis=0).astype(np.complex128)
    real_b = np.stack([item["real_B_valid"] for item in results], axis=0).astype(np.complex128)
    if theta_a.shape != (STATE_COUNT, 19) or theta_c.shape != (STATE_COUNT, 19):
        raise RuntimeError("Envelope19 coefficient shape contract failed")
    np.savez(
        RESULT_ROOT / "04_envelope19_aend_c2_coefficients.npz",
        state_ids=np.arange(STATE_COUNT, dtype=np.int64),
        theta_aend=theta_a,
        theta_c2=theta_c,
        basis_ids=np.asarray(support_ids),
        support_hash=np.asarray(support_digest),
        K=np.asarray(19),
        dmax=np.asarray(2),
        lambda_value=np.asarray(0.0),
    )
    np.save(RESULT_ROOT / "05a_lut_fingerprints.npy", (common_phi @ theta_a.T).T.astype(np.complex128))
    query_fingerprints = (common_phi @ theta_c.T).T.astype(np.complex128)
    lut_fingerprints = np.load(RESULT_ROOT / "05a_lut_fingerprints.npy")
    np.save(RESULT_ROOT / "05b_query_fingerprints.npy", query_fingerprints)
    distance = compute_cnmse_matrix(query_fingerprints, lut_fingerprints)
    if distance.shape != (STATE_COUNT, STATE_COUNT):
        raise RuntimeError(f"fingerprint distance matrix shape is {distance.shape}")
    np.save(RESULT_ROOT / "05_envelope19_fingerprint_cnmse_matrix.npy", distance)
    _checkpoint(
        phase="fingerprints_complete",
        support_hash=support_digest,
        common_B_hash=EXPECTED_COMMON_B_SHA,
        raw_manifest_before=raw_before,
        raw_manifest_after=None,
        completed_model_state_ids=list(range(STATE_COUNT)),
        completed_retrieval_rows=0,
        coefficients_completed=True,
        fingerprints_completed=True,
        retrieval_completed=False,
        excel_completed=False,
        plot_completed=False,
    )

    selected, selected_distance, true_rank, tie_count = top1_retrieval(distance)
    real_distance = compute_cnmse_matrix(real_b, real_b)
    retrieved_real_b = real_distance[np.arange(STATE_COUNT), selected]
    metrics = model_frame.copy()
    metrics["State_n_Q"] = selected.astype(int)
    metrics["retrieval_fingerprint_CNMSE_dB"] = selected_distance
    metrics["retrieved_real_B_CNMSE_dB"] = retrieved_real_b
    metrics["State_index_delta"] = selected.astype(int) - np.arange(STATE_COUNT)
    metrics["State_index_abs_delta"] = np.abs(metrics["State_index_delta"].to_numpy(dtype=int))
    metrics["is_exact_state_hit"] = selected == np.arange(STATE_COUNT)
    metrics["retrieved_real_B_pass"] = retrieved_real_b < THRESHOLD_DB
    metrics.to_csv(RESULT_ROOT / "06_retrieval_results_all425.csv", index=False)
    diagnostics = metrics[
        [
            "State_n_R",
            "State_n_Q",
            "retrieval_fingerprint_CNMSE_dB",
            "State_index_delta",
            "State_index_abs_delta",
            "is_exact_state_hit",
            "retrieved_real_B_CNMSE_dB",
            "retrieved_real_B_pass",
            "Y_Aend_train_NMSE_dB",
            "Y_Aend_B_NMSE_dB",
            "Y_C2_train_NMSE_dB",
            "Y_C2_B_NMSE_dB",
        ]
    ].copy()
    diagnostics.to_csv(RESULT_ROOT / "07_retrieval_diagnostics_all425.csv", index=False)
    failed = metrics.loc[~metrics["retrieved_real_B_pass"]].copy()
    failed.to_csv(RESULT_ROOT / "08_failed_states.csv", index=False)

    e23 = pd.read_csv(E23_METRICS).sort_values("State_n_R").reset_index(drop=True)
    e18 = pd.read_csv(E18_METRICS).sort_values("State_n_R").reset_index(drop=True)
    if len(e23) != STATE_COUNT or len(e18) != STATE_COUNT:
        raise RuntimeError("historical E18/E23 retrieval metrics are not canonical 425 rows")
    if int(e23["is_exact_state_hit"].sum()) != 123 or int(e23["retrieved_real_B_pass"].sum()) != 401:
        raise RuntimeError("historical E23 retrieval baseline changed")
    transition = np.where(
        e23["retrieved_real_B_pass"].to_numpy(dtype=bool) & metrics["retrieved_real_B_pass"].to_numpy(dtype=bool),
        "unchanged_pass",
        np.where(
            ~e23["retrieved_real_B_pass"].to_numpy(dtype=bool) & metrics["retrieved_real_B_pass"].to_numpy(dtype=bool),
            "E23_fail_to_E19_pass",
            np.where(
                e23["retrieved_real_B_pass"].to_numpy(dtype=bool) & ~metrics["retrieved_real_B_pass"].to_numpy(dtype=bool),
                "E23_pass_to_E19_fail",
                "unchanged_fail",
            ),
        ),
    )
    paired = pd.DataFrame(
        {
            "State_R": np.arange(STATE_COUNT),
            "E23_State_Q": e23["State_n_Q"].to_numpy(dtype=int),
            "E19_State_Q": selected,
            "E23_fingerprint_CNMSE_dB": e23["retrieval_fingerprint_CNMSE_dB"],
            "E19_fingerprint_CNMSE_dB": selected_distance,
            "E23_real_B_CNMSE_dB": e23["retrieved_real_B_CNMSE_dB"],
            "E19_real_B_CNMSE_dB": retrieved_real_b,
            "E23_pass": e23["retrieved_real_B_pass"],
            "E19_pass": metrics["retrieved_real_B_pass"],
            "transition": transition,
        }
    )
    paired.to_csv(RESULT_ROOT / "09_e23_vs_e19_paired_comparison.csv", index=False)

    quality_rows = _format_quality(model_frame, MODEL_NAME, "E19")
    quality_rows.extend(_format_quality(e23.rename(columns={
        "Y_Aend_train_NMSE_dB": "Y_Aend_train_NMSE_dB",
        "Y_Aend_B_NMSE_dB": "Y_Aend_B_NMSE_dB",
        "Y_C2_train_NMSE_dB": "Y_C2_train_NMSE_dB",
        "Y_C2_B_NMSE_dB": "Y_C2_B_NMSE_dB",
    }), "Envelope23-ilcEnd", "E23"))
    pd.DataFrame(quality_rows).to_csv(RESULT_ROOT / "14_model_quality_comparison.csv", index=False)
    _write_excel(metrics, RESULT_ROOT / "10_scenario_2_envelope19_c2endshared_full_lut_state_summary.xlsx")
    _write_plot(metrics, RESULT_ROOT / "11_envelope19_full_lut_retrieval.png")

    e19_summary = {
        "exact_count": int(metrics["is_exact_state_hit"].sum()),
        "real_B_pass_count": int(metrics["retrieved_real_B_pass"].sum()),
        "failure_count": int((~metrics["retrieved_real_B_pass"]).sum()),
        "nonself_count": int((metrics["State_n_Q"] != np.arange(STATE_COUNT)).sum()),
        "nonself_pass_count": int(((metrics["State_n_Q"] != np.arange(STATE_COUNT)) & metrics["retrieved_real_B_pass"]).sum()),
        "fingerprint_CNMSE_finite_only": _finite_summary(selected_distance),
        "real_B_CNMSE_finite_only": _finite_summary(retrieved_real_b),
        "exact_real_B_negative_infinity_count": int(np.isneginf(retrieved_real_b).sum()),
        "state_index_abs_delta_median": float(np.median(metrics["State_index_abs_delta"])),
        "state_index_abs_delta_Q90": float(np.quantile(metrics["State_index_abs_delta"], 0.90)),
        "state_index_abs_delta_max": int(np.max(metrics["State_index_abs_delta"])),
    }
    classes = paired["transition"].value_counts().to_dict()
    transitions = {key: int(classes.get(key, 0)) for key in ("unchanged_pass", "E23_fail_to_E19_pass", "E23_pass_to_E19_fail", "unchanged_fail")}
    recovered = paired.loc[paired["transition"] == "E23_fail_to_E19_pass", "State_R"].astype(int).tolist()
    regressed = paired.loc[paired["transition"] == "E23_pass_to_E19_fail", "State_R"].astype(int).tolist()
    quality_json = pd.DataFrame(quality_rows).to_dict("records")
    summary_lines = [
        f"Task: {TASK_NAME}",
        "Mode: evaluation_only; Frozen Envelope19-C2EndShared; only support changed versus E23 Full-LUT protocol.",
        f"K=19; dmax=2; lambda=0; solver=OLS; support_hash={support_digest}",
        "ABC bounds: A=[0,12288), B=[12288,17203), C=[17203,24576).",
        "allow_self=true; common-B uses canonical State 0 and the frozen E23 hash.",
        "",
        f"E19 vs E23 model-quality comparison: {json.dumps(quality_json, ensure_ascii=False, sort_keys=True)}",
        f"E19 retrieval summary: {json.dumps(e19_summary, ensure_ascii=False, sort_keys=True)}",
        f"E19 transitions versus E23: {json.dumps(transitions, ensure_ascii=False, sort_keys=True)}",
        f"Recovered State IDs: {recovered}",
        f"Regressed State IDs: {regressed}",
        f"Historical E18 exact/Real-B pass: {int(e18['is_exact_state_hit'].sum())}/425, {int(e18['retrieved_real_B_pass'].sum())}/425.",
        f"Historical E23 exact/Real-B pass: {int(e23['is_exact_state_hit'].sum())}/425, {int(e23['retrieved_real_B_pass'].sum())}/425.",
        "Real-B pass is the primary retrieval criterion; fingerprint CNMSE is a ranking diagnostic.",
        "No support selection, Ridge, dmax scan, Top-k fallback, clustering, Type-III, DPD replay, or low-bandwidth task was run.",
        f"Raw snapshot before: {json.dumps(raw_before, ensure_ascii=False, sort_keys=True)}",
        f"Raw snapshot after: {json.dumps(raw_manifest_gate(), ensure_ascii=False, sort_keys=True)}",
        f"raw_data_modified: {raw_before != raw_manifest_gate()}",
    ]
    (RESULT_ROOT / "12_final_result_summary.txt").write_text("\n".join(summary_lines) + "\n", encoding="utf-8")
    raw_after = raw_manifest_gate()
    _checkpoint(
        phase="completed",
        support_hash=support_digest,
        common_B_hash=EXPECTED_COMMON_B_SHA,
        raw_manifest_before=raw_before,
        raw_manifest_after=raw_after,
        completed_model_state_ids=list(range(STATE_COUNT)),
        completed_retrieval_rows=STATE_COUNT,
        coefficients_completed=True,
        fingerprints_completed=True,
        retrieval_completed=True,
        excel_completed=True,
        plot_completed=True,
    )
    return {
        "task": TASK_NAME,
        "status": "SUCCESS",
        "envelope19_exact_state_hit_count": e19_summary["exact_count"],
        "envelope19_real_B_pass_count": e19_summary["real_B_pass_count"],
        "envelope19_failure_count": e19_summary["failure_count"],
        "e23_real_B_pass_count": int(e23["retrieved_real_B_pass"].sum()),
        "recovered": transitions["E23_fail_to_E19_pass"],
        "regressed": transitions["E23_pass_to_E19_fail"],
        "raw_data_modified": raw_before != raw_after,
        "result_root": str(RESULT_ROOT),
    }


def main() -> None:
    try:
        result = _run(resume="--resume" in sys.argv[1:])
    except Exception as exc:
        _append_log(WORK_LOG, f"[{_now()}] FAILED {TASK_NAME}: {type(exc).__name__}: {exc}\n")
        raise
    _append_log(
        WORK_LOG,
        f"[{_now()}] Completed {TASK_NAME}\n"
        f"Result: {json.dumps(result, ensure_ascii=False, sort_keys=True)}\n"
        "No re-selection, Ridge scan, Top-k fallback, clustering, Type-III, DPD, or low-bandwidth task was run.\n",
    )
    _append_log(
        HANDOFF_LOG,
        f"\n{_now()} | 新任务：{TASK_NAME}\n"
        f"完成 Frozen Envelope19-C2EndShared 的 5B Full425 common-B LUT retrieval；结果目录：{RESULT_ROOT}\n"
        f"结果摘要：{json.dumps(result, ensure_ascii=False, sort_keys=True)}\n"
        "完整复用 E23 Full-LUT protocol，仅替换 support；未运行新的 support/Ridge/聚类/DPD/low-bandwidth。\n",
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))


if __name__ == "__main__":
    __import__("multiprocessing").freeze_support()
    main()
